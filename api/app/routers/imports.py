"""Importador CSV/Excel (migracion desde Fygaro u hojas propias).

Flujo en dos pasos: primero `commit=false` devuelve la vista previa fila por fila (nuevo / actualizar / error),
luego `commit=true` aplica. Los encabezados se reconocen por sinonimos, asi que funcionan los exportes de Fygaro
("Nombre", "Identificacion", "Correo electronico"...) y plantillas propias. Clientes se casan por identificacion
(o correo), productos por codigo, proveedores por cedula (o nombre), facturas historicas por numero.
"""

from __future__ import annotations

import re
from datetime import date
from decimal import Decimal

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..core.db import get_db
from ..core.deps import Principal, require
from ..models import Category, Customer, Invoice, InvoiceLine, Product, ProductTax, Supplier, Tax
from ..services import tabular as tb
from ..services.documents import audit
from ..services.totals import d

router = APIRouter(prefix="/import", tags=["importador"])

MAX = 5 * 1024 * 1024

COLS = {
    "customers": {
        "name": ("nombre", "cliente", "razon_social", "nombre_cliente", "name", "nombre_o_razon_social", "nombre_completo"),
        "id_number": ("identificacion", "cedula", "numero_identificacion", "numero_de_identificacion", "id", "cedula_juridica", "cedula_fisica", "nit"),
        "id_type": ("tipo_identificacion", "tipo_de_identificacion", "tipo_id", "tipo"),
        "email": ("correo", "correo_electronico", "email", "e_mail", "correo_facturacion"),
        "phone": ("telefono", "tel", "celular", "phone", "movil"),
        "whatsapp": ("whatsapp", "wa"),
        "currency": ("moneda", "divisa", "currency"),
        "notes": ("notas", "observaciones", "comentarios"),
        "address": ("direccion", "otras_senas", "senas", "address"),
    },
    "products": {
        "code": ("codigo", "sku", "code", "codigo_interno", "codigo_producto", "referencia"),
        "name": ("nombre", "producto", "descripcion", "nombre_producto", "name", "detalle"),
        "price": ("precio", "precio_unitario", "precio_venta", "price", "monto", "precio_sin_impuesto"),
        "item_type": ("tipo", "tipo_item", "producto_o_servicio"),
        "cabys_code": ("cabys", "codigo_cabys", "cabys_code"),
        "tax_rate": ("iva", "impuesto", "tarifa", "tarifa_iva", "tax"),
        "unit": ("unidad", "unidad_medida", "unidad_de_medida", "unit"),
        "currency": ("moneda", "divisa", "currency"),
        "description_invoice": ("descripcion_larga", "descripcion_factura", "detalle_factura"),
        "stock_min": ("stock_minimo", "minimo", "min_stock"),
    },
    "suppliers": {
        "name": ("nombre", "proveedor", "razon_social", "name"),
        "tax_id": ("identificacion", "cedula", "cedula_juridica", "tax_id", "nit"),
        "email": ("correo", "correo_electronico", "email"),
        "phone": ("telefono", "tel", "celular", "phone"),
    },
    "invoices": {
        "number": ("numero", "numero_factura", "factura", "consecutivo", "no_factura", "number", "documento"),
        "date": ("fecha", "fecha_emision", "fecha_de_emision", "date"),
        "customer": ("cliente", "nombre_cliente", "receptor", "customer"),
        "customer_id_number": ("identificacion", "cedula", "identificacion_cliente", "cedula_cliente"),
        "subtotal": ("subtotal", "monto_sin_impuesto", "base", "total_venta_neta", "venta_neta"),
        "tax": ("impuesto", "iva", "total_impuesto", "monto_iva"),
        "total": ("total", "monto_total", "total_comprobante", "monto"),
        "balance": ("saldo", "pendiente", "por_cobrar", "saldo_pendiente", "balance"),
        "currency": ("moneda", "divisa", "currency"),
        "status": ("estado", "status"),
        "clave": ("clave", "clave_numerica"),
    },
}

TEMPLATES = {
    "customers": ["Nombre", "Tipo de identificacion", "Identificacion", "Correo electronico", "Telefono", "WhatsApp", "Moneda", "Notas"],
    "products": ["Codigo", "Nombre", "Precio", "Tipo", "CABYS", "IVA", "Unidad", "Moneda"],
    "suppliers": ["Nombre", "Identificacion", "Correo electronico", "Telefono"],
    "invoices": ["Numero", "Fecha", "Cliente", "Identificacion", "Subtotal", "Impuesto", "Total", "Saldo", "Moneda", "Clave"],
    "catalogo": ["Categoria", "Codigo", "Descripcion", "Stock", "Precio"],
}


def _row(kind: str, r: dict) -> dict:
    return {k: tb.pick(r, *aliases) for k, aliases in COLS[kind].items()}


def _id_type(raw, number: str | None) -> str:
    s = tb.norm(raw)
    if "jur" in s or s in ("02", "2"):
        return "juridica"
    if "dimex" in s or s in ("03", "3"):
        return "dimex"
    if "nite" in s or s in ("04", "4"):
        return "nite"
    if "extr" in s or "pasap" in s:
        return "extranjero"
    if "fis" in s or s in ("01", "1"):
        return "fisica"
    digits = "".join(ch for ch in (number or "") if ch.isdigit())
    return "juridica" if len(digits) == 10 and digits.startswith("3") else ("dimex" if len(digits) in (11, 12) else "fisica")


def _currency(v) -> str:
    s = tb.norm(v)
    return "USD" if s in ("usd", "dolares", "dolar", "us") else "CRC"


# ---------- procesadores: devuelven (accion, detalle) y aplican si commit ----------
def _customers(db: Session, tid: int, rows: list[dict], commit: bool):
    out = []
    seen = set()
    for i, r in enumerate(rows, start=2):
        v = _row("customers", r)
        name = tb.text(v["name"], 200)
        idn = tb.text(v["id_number"], 30)
        email = (tb.text(v["email"], 200) or "").lower() or None
        if not name:
            out.append({"row": i, "action": "error", "detail": "Falta el nombre"})
            continue
        key = idn or email or name.lower()
        if key in seen:
            out.append({"row": i, "action": "omitir", "detail": f"Duplicado en el archivo: {name}"})
            continue
        seen.add(key)
        c = None
        if idn:
            c = db.scalar(select(Customer).where(Customer.tenant_id == tid, Customer.id_number == idn))
        if not c and email:
            c = db.scalar(select(Customer).where(Customer.tenant_id == tid, func.lower(Customer.email) == email))
        action = "actualizar" if c else "nuevo"
        if commit:
            if not c:
                c = Customer(tenant_id=tid, name=name)
                db.add(c)
            c.name = name
            c.id_number = idn or c.id_number
            c.id_type = _id_type(v["id_type"], idn) if (v["id_type"] or idn) else c.id_type
            c.email = email or c.email
            c.phone = tb.text(v["phone"], 40) or c.phone
            c.whatsapp = tb.text(v["whatsapp"], 40) or c.whatsapp or c.phone
            c.currency = _currency(v["currency"]) if v["currency"] else (c.currency or "CRC")
            if v["notes"]:
                c.notes = tb.text(v["notes"], 2000)
            if v["address"]:
                c.address = {**(c.address or {}), "senas": tb.text(v["address"], 300)}
        out.append({"row": i, "action": action, "detail": f"{name} · {idn or email or 'sin identificación'}"})
    return out


def _products(db: Session, tid: int, rows: list[dict], commit: bool):
    taxes = {d(t.rate): t for t in db.scalars(select(Tax).where(Tax.tenant_id == tid, Tax.active))}
    out = []
    seen = set()
    for i, r in enumerate(rows, start=2):
        v = _row("products", r)
        code = tb.text(v["code"], 60)
        name = tb.text(v["name"], 200)
        price = tb.num(v["price"])
        if not code or not name:
            out.append({"row": i, "action": "error", "detail": "Faltan código o nombre"})
            continue
        if code in seen:
            out.append({"row": i, "action": "omitir", "detail": f"Código repetido en el archivo: {code}"})
            continue
        seen.add(code)
        if price is None or price < 0:
            out.append({"row": i, "action": "error", "detail": f"{code}: precio inválido"})
            continue
        rate = tb.num(v["tax_rate"])
        if rate is not None and rate < 1 and rate > 0:
            rate = rate * 100  # 0.13 -> 13
        rate = Decimal(13) if rate is None else rate.quantize(Decimal(1))
        if rate not in taxes:
            out.append({"row": i, "action": "error", "detail": f"{code}: tarifa de IVA {rate}% no configurada"})
            continue
        cabys = "".join(ch for ch in str(tb.text(v["cabys_code"]) or "") if ch.isdigit()) or None
        if cabys and len(cabys) != 13:
            out.append({"row": i, "action": "error", "detail": f"{code}: CABYS debe tener 13 dígitos"})
            continue
        kind = "servicio" if "serv" in tb.norm(v["item_type"]) else "producto"
        p = db.scalar(select(Product).where(Product.tenant_id == tid, Product.code == code))
        action = "actualizar" if p else "nuevo"
        if commit:
            if not p:
                p = Product(tenant_id=tid, code=code, name=name)
                db.add(p)
            p.name, p.price, p.item_type = name, price, kind
            p.cabys_code = cabys or p.cabys_code
            p.unit = tb.text(v["unit"], 10) or p.unit or ("Sp" if kind == "servicio" else "Unid")
            p.currency = _currency(v["currency"]) if v["currency"] else (p.currency or "CRC")
            if v["description_invoice"]:
                p.description_invoice = tb.text(v["description_invoice"], 2000)
            p.taxes.clear()
            db.flush()
            p.taxes.append(ProductTax(tax_id=taxes[rate].id))
        out.append({"row": i, "action": action, "detail": f"{code} · {name} · {price:,.2f} · IVA {rate}%"})
    return out


def _suppliers(db: Session, tid: int, rows: list[dict], commit: bool):
    out = []
    for i, r in enumerate(rows, start=2):
        v = _row("suppliers", r)
        name = tb.text(v["name"], 160)
        if not name:
            out.append({"row": i, "action": "error", "detail": "Falta el nombre"})
            continue
        tax_id = tb.text(v["tax_id"], 30)
        s = db.scalar(select(Supplier).where(Supplier.tenant_id == tid, Supplier.tax_id == tax_id)) if tax_id else None
        s = s or db.scalar(select(Supplier).where(Supplier.tenant_id == tid, func.lower(Supplier.name) == name.lower()))
        action = "actualizar" if s else "nuevo"
        if commit:
            if not s:
                s = Supplier(tenant_id=tid, name=name)
                db.add(s)
            s.name, s.tax_id = name, tax_id or s.tax_id
            s.email = tb.text(v["email"], 200) or s.email
            s.phone = tb.text(v["phone"], 40) or s.phone
        out.append({"row": i, "action": action, "detail": f"{name} · {tax_id or 'sin cédula'}"})
    return out


def _invoices(db: Session, tid: int, rows: list[dict], commit: bool):
    """Facturas historicas (ya emitidas en Fygaro): quedan como referencia para reportes, saldos y la ficha del
    cliente. No se reenvian a Hacienda ni consumen consecutivo (einvoice_status = importada)."""
    out = []
    for i, r in enumerate(rows, start=2):
        v = _row("invoices", r)
        number = tb.text(v["number"], 40)
        dt = tb.to_date(v["date"])
        total = tb.num(v["total"])
        if not number or not dt or total is None:
            out.append({"row": i, "action": "error", "detail": "Faltan número, fecha o total"})
            continue
        if db.scalar(select(Invoice).where(Invoice.tenant_id == tid, Invoice.number == number)):
            out.append({"row": i, "action": "omitir", "detail": f"{number} ya existe"})
            continue
        tax = tb.num(v["tax"]) or Decimal(0)
        sub = tb.num(v["subtotal"])
        sub = sub if sub is not None else total - tax
        bal = tb.num(v["balance"])
        st = tb.norm(v["status"])
        if bal is None:
            bal = Decimal(0) if ("pag" in st or not st) else total
        status = "anulada" if "anul" in st else ("pagada" if bal <= 0 else ("parcial" if bal < total else "creado"))
        cid_num = tb.text(v["customer_id_number"], 30)
        cust = db.scalar(select(Customer).where(Customer.tenant_id == tid, Customer.id_number == cid_num)) if cid_num else None
        cname = tb.text(v["customer"], 200)
        if not cust and cname:
            cust = db.scalar(select(Customer).where(Customer.tenant_id == tid, func.lower(Customer.name) == cname.lower()))
        detail = f"{number} · {dt.isoformat()} · {cname or '—'} · {total:,.2f}" + ("" if cust or not cname else " (cliente nuevo)")
        if commit:
            if not cust and cname:
                cust = Customer(tenant_id=tid, name=cname, id_number=cid_num, id_type=_id_type(None, cid_num))
                db.add(cust)
                db.flush()
            cur = _currency(v["currency"]) if v["currency"] else "CRC"
            rate = (tax / sub * 100).quantize(Decimal("0.001")) if sub else Decimal(0)
            inv = Invoice(
                tenant_id=tid,
                number=number,
                consecutive=None,
                doc_type="FE",
                customer_id=cust.id if cust else None,
                currency=cur,
                fx_sell=1,
                fx_buy=1,
                issue_date=dt,
                due_date=dt,
                subtotal=sub,
                discount_total=0,
                tax_total=tax,
                total=total,
                balance=max(bal, Decimal(0)),
                status=status,
                einvoice_status="importada",
                clave=tb.text(v["clave"], 50),
                internal_notes="Importada (historico de Fygaro)",
                sale_condition="01",
                credit_days=0,
                payment_method="01",
            )
            inv.lines.append(
                InvoiceLine(
                    position=1,
                    name="Venta registrada en sistema anterior",
                    unit="Sp",
                    quantity=1,
                    unit_price=sub,
                    tax_rate=rate,
                    subtotal=sub,
                    tax_amount=tax,
                    total=total,
                )
            )
            db.add(inv)
        out.append({"row": i, "action": "nuevo", "detail": detail})
    return out


# ---------- lista de precios de proveedor (Eurocomp / Hikvision y similares) ----------
CATALOG_COLS = {
    "code": ("codigo", "code", "sku", "parte", "referencia"),
    "name": ("descripcion", "producto", "nombre", "detalle", "description"),
    "cost": ("precio", "costo", "price", "cost", "precio_unitario", "precio_lista"),
    "stock": ("stock", "existencia", "existencias", "disponible", "inventario"),
    "section": ("fotografia", "categoria", "familia", "linea"),
}
GROUPS = (
    "ACCESORIOS",
    "AUDIO Y SONIDO",
    "DISPOSITIVOS DE VISUALIZACION",
    "REDES Y CONECTIVDAD",
    "REDES Y CONECTIVIDAD",
    "RESPALDO ELECTRICO",
    "RESPALDO ELÉCTRICO",
    "SEGURIDAD",
    "CCTV",
    "VIDEOVIGILANCIA",
)
MODEL_RE = re.compile(r"\b((?:i?DS|AE|HWI|HWT|DS)-[A-Z0-9][A-Z0-9\-/().]{2,})", re.I)


def _split_section(raw: str) -> tuple[str, str]:
    """ "SEGURIDAD Camaras" -> ("Seguridad", "Camaras"). Si no trae grupo conocido, todo es categoria."""
    t = " ".join(str(raw or "").split())
    up = tb.norm(t).replace("_", " ")
    for g in GROUPS:
        gn = tb.norm(g).replace("_", " ")
        if up.startswith(gn):
            sub = t[len(g) :].strip(" -·").strip()
            return t[: len(g)].strip().title(), (sub.title() if sub else t.title())
    return "", t.title()


def _category(db: Session, tid: int, group: str, name: str, cache: dict, commit: bool):
    if not name:
        return None
    key = (group, name)
    if key in cache:
        return cache[key]
    parent = None
    if group:
        parent = db.scalar(select(Category).where(Category.tenant_id == tid, Category.name == group, Category.parent_id.is_(None)))
        if not parent and commit:
            parent = Category(tenant_id=tid, name=group, show_on_web=True)
            db.add(parent)
            db.flush()
    cat = db.scalar(select(Category).where(Category.tenant_id == tid, Category.name == name))
    if not cat and commit:
        cat = Category(tenant_id=tid, name=name, parent_id=parent.id if parent else None, show_on_web=True)
        db.add(cat)
        db.flush()
    cache[key] = cat
    return cat


def _catalog(db: Session, tid: int, rows: list[dict], commit: bool, opts: dict | None = None):
    """Lista de precios de un proveedor: el costo entra tal cual y el precio de venta se calcula con el margen.
    Las filas sin codigo son encabezados de categoria (asi vienen los listados de Eurocomp)."""
    from datetime import UTC, datetime

    from ..models import Supplier, Tenant
    from ..services import pricing
    from ..services.documents import today_fx

    opts = opts or {}
    tenant = db.get(Tenant, tid)
    margin = d(opts.get("margin") if opts.get("margin") is not None else pricing.default_margin(tenant))
    cost_currency = (opts.get("currency") or "USD").upper()
    supplier_name = (opts.get("supplier") or "Proveedor").strip()
    default_brand = (opts.get("brand") or "").strip() or None
    fx = today_fx(db, "USD")[0] if cost_currency != "CRC" else Decimal(1)

    supplier = db.scalar(select(Supplier).where(Supplier.tenant_id == tid, func.lower(Supplier.name) == supplier_name.lower()))
    if not supplier and commit:
        supplier = Supplier(tenant_id=tid, name=supplier_name)
        db.add(supplier)
        db.flush()
    iva13 = db.scalar(select(Tax).where(Tax.tenant_id == tid, Tax.rate_code == "08"))
    cats: dict = {}
    section = ("", "")
    out, seen = [], set()
    for i, r in enumerate(rows, start=2):
        v = {k: tb.pick(r, *aliases) for k, aliases in CATALOG_COLS.items()}
        code = tb.text(v["code"], 60)
        if not code:
            if v["section"]:
                section = _split_section(str(v["section"]))
            continue
        raw_name = str(v["name"] or "").strip()
        if not raw_name:
            out.append({"row": i, "action": "error", "detail": f"{code}: sin descripción"})
            continue
        cost = tb.num(v["cost"])
        if cost is None or cost < 0:
            out.append({"row": i, "action": "error", "detail": f"{code}: costo inválido"})
            continue
        if code in seen:
            out.append({"row": i, "action": "omitir", "detail": f"Código repetido en el archivo: {code}"})
            continue
        seen.add(code)
        parts = [x.strip() for x in raw_name.splitlines() if x.strip()]
        headline = parts[0][:200]
        specs = "\n".join(parts[1:])[:4000]
        model = MODEL_RE.search(headline) or MODEL_RE.search(raw_name)
        model = model.group(1).upper() if model else None
        brand = "Hikvision" if "hikvision" in raw_name.lower() else default_brand
        stock = tb.num(v["stock"])
        price = pricing.sale_price(db, cost, cost_currency, "CRC", margin, fx)
        pr = db.scalar(select(Product).where(Product.tenant_id == tid, Product.code == code))
        action = "actualizar" if pr else "nuevo"
        group, cat_name = section
        if commit:
            cat = _category(db, tid, group, cat_name, cats, commit)
            if not pr:
                pr = Product(tenant_id=tid, code=code, name=headline, item_type="producto", unit="Unid")
                db.add(pr)
                db.flush()
                if iva13:
                    pr.taxes.append(ProductTax(tax_id=iva13.id))
            pr.name = headline
            pr.description_invoice = headline
            pr.description_store = specs or pr.description_store
            pr.cost, pr.cost_currency, pr.margin_pct = cost, cost_currency, margin
            pr.price, pr.currency = price, "CRC"
            pr.brand, pr.model = brand, model or pr.model
            pr.supplier_id = supplier.id if supplier else pr.supplier_id
            pr.supplier_sku = code
            pr.supplier_stock = int(stock) if stock is not None else None
            pr.supplier_updated_at = datetime.now(UTC)
            pr.category_id = cat.id if cat else pr.category_id
        out.append(
            {
                "row": i,
                "action": action,
                "detail": f"{code} · {headline[:60]} · costo {cost_currency} {cost:,.2f} → venta ₡{price:,.0f}" + (f" · {cat_name}" if cat_name else ""),
            }
        )
    return out


PROCESSORS = {"customers": _customers, "products": _products, "suppliers": _suppliers, "invoices": _invoices, "catalogo": _catalog}
PERM = {
    "customers": ("crm", "crear"),
    "products": ("catalog", "crear"),
    "suppliers": ("catalog", "crear"),
    "invoices": ("sales", "crear"),
    "catalogo": ("catalog", "crear"),
}


@router.get("/{kind}/template")
def template(kind: str, p: Principal = Depends(require("dashboard", "ver"))):
    if kind not in TEMPLATES:
        raise HTTPException(404, "Tipo no soportado")
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = kind
    ws.append(TEMPLATES[kind])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="E2233A")
        ws.column_dimensions[c.column_letter].width = max(14, len(str(c.value)) + 4)
    buf = BytesIO()
    wb.save(buf)
    return Response(
        buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="plantilla-{kind}.xlsx"'},
    )


@router.post("/{kind}")
async def run_import(
    kind: str,
    commit: bool = False,
    margin: Decimal | None = None,
    supplier: str | None = None,
    brand: str | None = None,
    currency: str = "USD",
    file: UploadFile = File(...),
    p: Principal = Depends(require("dashboard", "ver")),
    db: Session = Depends(get_db),
):
    if kind not in PROCESSORS:
        raise HTTPException(404, "Tipo no soportado")
    if kind == "catalogo" and margin is not None and not (0 <= margin <= 95):
        raise HTTPException(422, "El margen debe estar entre 0 y 95 %")
    mod, act = PERM[kind]
    if not p.can(mod, act):
        raise HTTPException(403, f"Sin permiso: {mod}.{act}")
    data = await file.read(MAX + 1)
    if len(data) > MAX:
        raise HTTPException(413, "Archivo muy grande (maximo 5 MB)")
    try:
        rows = tb.read_table(data, file.filename or "")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el archivo: {e}") from e
    if not rows:
        raise HTTPException(422, "El archivo no tiene filas")
    if len(rows) > 5000:
        raise HTTPException(422, "Maximo 5000 filas por archivo; divídalo en partes")
    headers = sorted({h for r in rows for h in r})
    cols = CATALOG_COLS if kind == "catalogo" else COLS[kind]
    recognized = {field: next((a for a in aliases if a in headers), None) for field, aliases in cols.items()}
    opts = {"margin": margin, "supplier": supplier, "brand": brand, "currency": currency}
    result = PROCESSORS[kind](db, p.tenant.id, rows, commit, opts) if kind == "catalogo" else PROCESSORS[kind](db, p.tenant.id, rows, commit)
    summary = {a: sum(1 for r in result if r["action"] == a) for a in ("nuevo", "actualizar", "omitir", "error")}
    if commit:
        audit(db, p.tenant.id, p.user.id, "import", kind, None, {"file": file.filename, **summary}, ip=p.ip)
        db.commit()
    else:
        db.rollback()
    return {"kind": kind, "commit": commit, "rows": len(rows), "summary": summary, "columns": recognized, "result": result[:1000], "at": date.today()}
