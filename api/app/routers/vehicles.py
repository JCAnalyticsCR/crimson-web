"""Vehiculos y combustible (seccion 12 / AC21): flota, cargas, importacion con vista previa, reportes y Excel.

Permisos: modulo "vehicles". Admin y contabilidad ver/crear/editar/exportar; supervisor solo ver.
No hay borrado de cargas: lo que no debe contar se EXCLUYE con motivo y sigue visible en la lista.
"""

from __future__ import annotations

from datetime import date, timedelta
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from ..core.db import get_db
from ..core.deps import Principal, require
from ..models import FuelLoad, Media, Project, Tenant, Vehicle, VehicleCost
from ..services import tabular as tb
from ..services import vehicles as svc
from ..services.documents import audit

router = APIRouter(prefix="/vehicles", tags=["vehiculos"])

MAX = 5 * 1024 * 1024
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------- esquemas ----------
class VehicleIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    plate: str | None = Field(None, max_length=20)
    kind: str | None = Field(None, max_length=30)
    owner: Literal["empresa", "personal"] = "empresa"
    driver: str | None = Field(None, max_length=120)
    active: bool = True
    excluded_from_analysis: bool = False
    exclusion_reason: str | None = Field(None, max_length=240)
    notes: str | None = Field(None, max_length=4000)

    @field_validator("name", "plate", "kind", "driver", "exclusion_reason", "notes")
    @classmethod
    def _strip(cls, v):
        return v.strip() or None if isinstance(v, str) else v


class FuelIn(BaseModel):
    vehicle_id: int
    date: date
    driver: str | None = Field(None, max_length=120)
    liters: Decimal | None = Field(None, gt=0, le=2000)
    amount: Decimal = Field(gt=0, le=Decimal("100000000"))
    currency: Literal["CRC", "USD"] = "CRC"
    odometer: int | None = Field(None, ge=0, le=5_000_000)
    full_tank: bool = False
    receipt_media_id: int | None = None
    usage: str | None = Field(None, max_length=160)
    project_id: int | None = None
    excluded: bool = False
    exclusion_reason: str | None = Field(None, max_length=240)
    notes: str | None = Field(None, max_length=4000)


class CostIn(BaseModel):
    vehicle_id: int
    date: date
    kind: Literal["mantenimiento", "llantas", "seguro", "depreciacion", "otro"]
    amount: Decimal = Field(gt=0, le=Decimal("1000000000"))
    currency: Literal["CRC", "USD"] = "CRC"
    description: str | None = Field(None, max_length=240)


class ConfigIn(BaseModel):
    monthly_reference: Decimal = Field(ge=0, le=Decimal("1000000000"))


# ---------- helpers ----------
def _vehicle(db: Session, tid: int, vid: int) -> Vehicle:
    v = db.get(Vehicle, vid)
    if not v or v.tenant_id != tid:
        raise HTTPException(404, "Vehículo no encontrado")
    return v


def _vout(v: Vehicle, counts: dict | None = None) -> dict:
    c = (counts or {}).get(v.id, {})
    return {
        "id": v.id,
        "name": v.name,
        "plate": v.plate,
        "kind": v.kind,
        "owner": v.owner,
        "driver": v.driver,
        "active": v.active,
        "excluded_from_analysis": v.excluded_from_analysis,
        "exclusion_reason": v.exclusion_reason,
        "notes": v.notes,
        "loads": c.get("n", 0),
        "last_date": c.get("last"),
        "last_odometer": c.get("odo"),
    }


def _check_refs(db: Session, tid: int, data: FuelIn) -> None:
    _vehicle(db, tid, data.vehicle_id)
    if data.project_id:
        pr = db.get(Project, data.project_id)
        if not pr or pr.tenant_id != tid:
            raise HTTPException(422, "Proyecto no encontrado")
    if data.receipt_media_id:
        m = db.get(Media, data.receipt_media_id)
        if not m or m.tenant_id != tid:
            raise HTTPException(422, "Comprobante no encontrado")
    if data.excluded and not data.exclusion_reason:
        raise HTTPException(422, "Indique el motivo de la exclusión")


def _lout(x: FuelLoad, vehicles: dict[int, Vehicle], projects: dict[int, Project], media: dict[int, Media]) -> dict:
    from .media import public_url

    v = vehicles.get(x.vehicle_id)
    why = svc.exclusion(x, v)
    m = media.get(x.receipt_media_id) if x.receipt_media_id else None
    pr = projects.get(x.project_id) if x.project_id else None
    return {
        "id": x.id,
        "vehicle_id": x.vehicle_id,
        "vehicle": v.name if v else "—",
        "plate": v.plate if v else None,
        "date": x.date.isoformat(),
        "driver": x.driver,
        "liters": x.liters,
        "amount": x.amount,
        "currency": x.currency,
        "price_per_liter": (Decimal(str(x.amount)) / Decimal(str(x.liters))).quantize(Decimal("0.01")) if x.liters else None,
        "odometer": x.odometer,
        "full_tank": x.full_tank,
        "receipt_media_id": x.receipt_media_id,
        "receipt_url": public_url(m) if m else None,
        "receipt_name": m.filename if m else None,
        "usage": x.usage,
        "project_id": x.project_id,
        "project": f"{pr.number} · {pr.name}" if pr else None,
        "excluded": x.excluded,
        "exclusion_reason": x.exclusion_reason,
        "excluded_effective": why is not None,
        "exclusion_effective": why,
        "source": x.source,
        "notes": x.notes,
    }


def _period(frm: date | None, to: date | None) -> tuple[date, date]:
    b = to or date.today()
    a = frm or date(b.year, 1, 1)
    if a > b:
        raise HTTPException(422, "La fecha inicial es posterior a la final")
    if (b - a).days > 3700:
        raise HTTPException(422, "Periodo máximo: 10 años")
    return a, b


def _loads(db: Session, tid: int, a: date | None, b: date | None, vehicle_id: int | None, driver: str | None, status: str, q: str | None, limit: int = 2000) -> list[dict]:
    qy = select(FuelLoad).where(FuelLoad.tenant_id == tid)
    if a:
        qy = qy.where(FuelLoad.date >= a)
    if b:
        qy = qy.where(FuelLoad.date <= b)
    if vehicle_id:
        qy = qy.where(FuelLoad.vehicle_id == vehicle_id)
    if driver:
        qy = qy.where(FuelLoad.driver.ilike(f"%{driver.strip()[:60]}%"))
    if q:
        like = f"%{q.strip()[:60]}%"
        qy = qy.where(or_(FuelLoad.usage.ilike(like), FuelLoad.notes.ilike(like), FuelLoad.driver.ilike(like)))
    rows = db.scalars(qy.order_by(FuelLoad.date.desc(), FuelLoad.id.desc()).limit(limit)).all()
    vehicles = {v.id: v for v in db.scalars(select(Vehicle).where(Vehicle.tenant_id == tid)).all()}
    pids = {x.project_id for x in rows if x.project_id}
    mids = {x.receipt_media_id for x in rows if x.receipt_media_id}
    projects = {p.id: p for p in db.scalars(select(Project).where(Project.id.in_(pids)))} if pids else {}
    media = {m.id: m for m in db.scalars(select(Media).where(Media.id.in_(mids)))} if mids else {}
    out = [_lout(x, vehicles, projects, media) for x in rows]
    if status == "incluidas":
        out = [x for x in out if not x["excluded_effective"]]
    elif status == "excluidas":
        out = [x for x in out if x["excluded_effective"]]
    return out


# ---------- vehiculos ----------
@router.get("")
def list_vehicles(p: Principal = Depends(require("vehicles", "ver")), db: Session = Depends(get_db)):
    vs = db.scalars(select(Vehicle).where(Vehicle.tenant_id == p.tenant.id).order_by(Vehicle.active.desc(), Vehicle.name)).all()
    counts: dict[int, dict] = {}
    for x in db.scalars(select(FuelLoad).where(FuelLoad.tenant_id == p.tenant.id).order_by(FuelLoad.date, FuelLoad.id)).all():
        c = counts.setdefault(x.vehicle_id, {"n": 0, "last": None, "odo": None})
        c["n"] += 1
        c["last"] = x.date.isoformat()
        if x.odometer is not None:
            c["odo"] = x.odometer
    return [_vout(v, counts) for v in vs]


@router.post("", status_code=201)
def create_vehicle(data: VehicleIn, p: Principal = Depends(require("vehicles", "crear")), db: Session = Depends(get_db)):
    if data.excluded_from_analysis and not data.exclusion_reason:
        raise HTTPException(422, "Indique el motivo de la exclusión")
    v = Vehicle(tenant_id=p.tenant.id, **data.model_dump())
    db.add(v)
    db.flush()
    audit(db, p.tenant.id, p.user.id, "create", "vehicle", v.id, {"name": v.name, "plate": v.plate}, ip=p.ip)
    db.commit()
    return _vout(v)


@router.put("/{vid:int}")
def update_vehicle(vid: int, data: VehicleIn, p: Principal = Depends(require("vehicles", "editar")), db: Session = Depends(get_db)):
    v = _vehicle(db, p.tenant.id, vid)
    if data.excluded_from_analysis and not data.exclusion_reason:
        raise HTTPException(422, "Indique el motivo de la exclusión")
    before = {"excluded": v.excluded_from_analysis, "active": v.active}
    for k, val in data.model_dump().items():
        setattr(v, k, val)
    if not v.excluded_from_analysis:
        v.exclusion_reason = None
    audit(db, p.tenant.id, p.user.id, "update", "vehicle", v.id, {"before": before, "after": {"excluded": v.excluded_from_analysis, "active": v.active, "reason": v.exclusion_reason}}, ip=p.ip)
    db.commit()
    return _vout(v)


@router.get("/lookups")
def lookups(p: Principal = Depends(require("vehicles", "ver")), db: Session = Depends(get_db)):
    """Vehiculos y proyectos (numero + nombre) para los selectores del formulario."""
    vs = db.scalars(select(Vehicle).where(Vehicle.tenant_id == p.tenant.id).order_by(Vehicle.name)).all()
    prs = db.scalars(
        select(Project).where(Project.tenant_id == p.tenant.id, Project.trashed_at.is_(None), Project.archived_at.is_(None)).order_by(Project.id.desc()).limit(300)
    ).all()
    return {
        "vehicles": [{"id": v.id, "name": v.name, "plate": v.plate, "driver": v.driver, "active": v.active, "excluded": v.excluded_from_analysis} for v in vs],
        "projects": [{"id": x.id, "label": f"{x.number} · {x.name}"} for x in prs],
        "config": svc.settings_of(p.tenant),
    }


@router.get("/config")
def get_config(p: Principal = Depends(require("vehicles", "ver"))):
    return {**svc.settings_of(p.tenant), "label": svc.REFERENCE_LABEL, "note": svc.NOTE}


@router.put("/config")
def put_config(data: ConfigIn, p: Principal = Depends(require("vehicles", "editar")), db: Session = Depends(get_db)):
    t = db.get(Tenant, p.tenant.id)
    s = dict(t.settings or {})
    veh = dict(s.get("vehicles") or {})
    old = veh.get("monthly_reference")
    veh["monthly_reference"] = str(data.monthly_reference)
    s["vehicles"] = veh
    t.settings = s  # reasignar: el JSON no detecta cambios internos
    audit(db, p.tenant.id, p.user.id, "update", "vehicle_config", None, {"monthly_reference": [old, str(data.monthly_reference)]}, ip=p.ip)
    db.commit()
    return {**svc.settings_of(t), "label": svc.REFERENCE_LABEL, "note": svc.NOTE}


# ---------- cargas ----------
@router.get("/fuel")
def list_fuel(
    frm: date | None = Query(None, alias="from"),
    to: date | None = None,
    vehicle_id: int | None = None,
    driver: str | None = None,
    status: Literal["todas", "incluidas", "excluidas"] = "todas",
    q: str | None = None,
    p: Principal = Depends(require("vehicles", "ver")),
    db: Session = Depends(get_db),
):
    return _loads(db, p.tenant.id, frm, to, vehicle_id, driver, status, q)


@router.post("/fuel", status_code=201)
def create_fuel(data: FuelIn, p: Principal = Depends(require("vehicles", "crear")), db: Session = Depends(get_db)):
    _check_refs(db, p.tenant.id, data)
    if data.date > date.today() + timedelta(days=1):
        raise HTTPException(422, "La fecha de la carga está en el futuro")
    x = FuelLoad(tenant_id=p.tenant.id, source="manual", created_by=p.user.id, **data.model_dump())
    if not x.driver:
        x.driver = _vehicle(db, p.tenant.id, data.vehicle_id).driver
    db.add(x)
    db.flush()
    audit(db, p.tenant.id, p.user.id, "create", "fuel_load", x.id, {"vehicle_id": x.vehicle_id, "date": x.date.isoformat(), "amount": str(x.amount)}, ip=p.ip)
    db.commit()
    return {"id": x.id}


@router.put("/fuel/{fid}")
def update_fuel(fid: int, data: FuelIn, p: Principal = Depends(require("vehicles", "editar")), db: Session = Depends(get_db)):
    x = db.get(FuelLoad, fid)
    if not x or x.tenant_id != p.tenant.id:
        raise HTTPException(404, "Carga no encontrada")
    _check_refs(db, p.tenant.id, data)
    before = {"excluded": x.excluded, "amount": str(x.amount), "odometer": x.odometer, "date": x.date.isoformat()}
    for k, val in data.model_dump().items():
        setattr(x, k, val)
    if not x.excluded:
        x.exclusion_reason = None
    audit(db, p.tenant.id, p.user.id, "update", "fuel_load", x.id, {"before": before, "after": {"excluded": x.excluded, "amount": str(x.amount), "odometer": x.odometer, "reason": x.exclusion_reason}}, ip=p.ip)
    db.commit()
    return {"id": x.id}


@router.get("/fuel/import/template")
def import_template(p: Principal = Depends(require("vehicles", "crear"))):
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "cargas"
    ws.append(svc.TEMPLATE)
    ws.append(["15/08/2026", "Pickup 1", "Conductor de ejemplo", 42.5, 35000, "CRC", 120500, "si", "Visita a cliente", "", ""])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="E2233A")
        ws.column_dimensions[c.column_letter].width = max(14, len(str(c.value)) + 4)
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": 'attachment; filename="plantilla-cargas-combustible.xlsx"'})


@router.post("/fuel/import")
async def import_fuel(
    commit: bool = False,
    create_vehicles: bool = False,
    file: UploadFile = File(...),
    p: Principal = Depends(require("vehicles", "crear")),
    db: Session = Depends(get_db),
):
    data = await file.read(MAX + 1)
    if len(data) > MAX:
        raise HTTPException(413, "Archivo muy grande (máximo 5 MB)")
    try:
        rows = tb.read_table(data, file.filename or "")
    except Exception as e:  # noqa: BLE001
        raise HTTPException(422, f"No se pudo leer el archivo: {e}") from e
    if not rows:
        raise HTTPException(422, "El archivo no tiene filas")
    if len(rows) > 5000:
        raise HTTPException(422, "Máximo 5000 filas por archivo; divídalo en partes")
    headers = sorted({h for r in rows for h in r})
    recognized = {k: next((a for a in al if a in headers), None) for k, al in svc.COLS.items()}
    result = svc.import_rows(db, p.tenant.id, rows, commit, p.user.id, create_vehicles)
    summary = {a: sum(1 for r in result if r["action"] == a) for a in ("nuevo", "omitir", "error")}
    summary["avisos"] = sum(1 for r in result if r["warnings"])
    if commit:
        audit(db, p.tenant.id, p.user.id, "import", "fuel_load", None, {"file": file.filename, **summary}, ip=p.ip)
        db.commit()
    else:
        db.rollback()
    return {"commit": commit, "rows": len(rows), "summary": summary, "columns": recognized, "result": result[:2000]}


# ---------- otros costos ----------
@router.get("/costs")
def list_costs(vehicle_id: int | None = None, p: Principal = Depends(require("vehicles", "ver")), db: Session = Depends(get_db)):
    qy = select(VehicleCost).where(VehicleCost.tenant_id == p.tenant.id)
    if vehicle_id:
        qy = qy.where(VehicleCost.vehicle_id == vehicle_id)
    names = {v.id: v.name for v in db.scalars(select(Vehicle).where(Vehicle.tenant_id == p.tenant.id))}
    return [
        {"id": c.id, "vehicle_id": c.vehicle_id, "vehicle": names.get(c.vehicle_id, "—"), "date": c.date.isoformat(), "kind": c.kind, "amount": c.amount, "currency": c.currency, "description": c.description}
        for c in db.scalars(qy.order_by(VehicleCost.date.desc(), VehicleCost.id.desc()).limit(2000))
    ]


@router.post("/costs", status_code=201)
def create_cost(data: CostIn, p: Principal = Depends(require("vehicles", "crear")), db: Session = Depends(get_db)):
    _vehicle(db, p.tenant.id, data.vehicle_id)
    c = VehicleCost(tenant_id=p.tenant.id, created_by=p.user.id, **data.model_dump())
    db.add(c)
    db.flush()
    audit(db, p.tenant.id, p.user.id, "create", "vehicle_cost", c.id, {"vehicle_id": c.vehicle_id, "kind": c.kind, "amount": str(c.amount)}, ip=p.ip)
    db.commit()
    return {"id": c.id}


@router.put("/costs/{cid}")
def update_cost(cid: int, data: CostIn, p: Principal = Depends(require("vehicles", "editar")), db: Session = Depends(get_db)):
    c = db.get(VehicleCost, cid)
    if not c or c.tenant_id != p.tenant.id:
        raise HTTPException(404, "Costo no encontrado")
    _vehicle(db, p.tenant.id, data.vehicle_id)
    before = {"amount": str(c.amount), "kind": c.kind, "date": c.date.isoformat()}
    for k, val in data.model_dump().items():
        setattr(c, k, val)
    audit(db, p.tenant.id, p.user.id, "update", "vehicle_cost", c.id, {"before": before}, ip=p.ip)
    db.commit()
    return {"id": c.id}


# ---------- reportes ----------
@router.get("/report")
def vehicle_report(
    frm: date | None = Query(None, alias="from"),
    to: date | None = None,
    vehicle_id: int | None = None,
    format: Literal["json", "xlsx"] = "json",
    p: Principal = Depends(require("vehicles", "ver")),
    db: Session = Depends(get_db),
):
    a, b = _period(frm, to)
    rep = svc.report(db, p.tenant, a, b, vehicle_id)
    if format == "xlsx":
        if not p.can("vehicles", "exportar"):
            raise HTTPException(403, "Sin permiso: vehicles.exportar")
        loads = _loads(db, p.tenant.id, a, b, vehicle_id, None, "todas", None, limit=20000)
        audit(db, p.tenant.id, p.user.id, "export", "vehicle_report", None, {"from": a.isoformat(), "to": b.isoformat()}, ip=p.ip)
        db.commit()
        return Response(
            svc.to_xlsx(rep, list(reversed(loads)), p.tenant.name),
            media_type=XLSX,
            headers={"Content-Disposition": f'attachment; filename="vehiculos-{a.isoformat()}-{b.isoformat()}.xlsx"'},
        )
    return rep
