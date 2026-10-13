"""Vehiculos y combustible: analisis de compras, km validos, km/L, costo por km e importacion de cargas.

Reglas (documento de Andres, seccion 12 / AC21):
- Las cargas excluidas, o de un vehiculo excluido, NO entran al analisis; se listan aparte con su marca.
- km/L solo con intervalos entre dos tanques llenos con odometro coherente (sube, sin cargas raras en medio).
- Comprar en un mes no equivale a consumir en ese mes: el resumen mensual es de COMPRAS, y los meses que el
  periodo o los datos cubren a medias se marcan como parciales.
- La "referencia historica mensual" es un parametro que alguien dijo de memoria: se muestra como referencia
  no validada, nunca como un promedio calculado.
"""

from __future__ import annotations

import calendar
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import ExchangeRate, FuelLoad, Project, Vehicle, VehicleCost
from . import tabular as tb

NOTE = "Comprar en un mes no equivale a consumir en ese mes."
REFERENCE_DEFAULT = Decimal(200000)
REFERENCE_LABEL = "Referencia histórica mensual (no validada)"
COST_KINDS = ("mantenimiento", "llantas", "seguro", "depreciacion", "otro")
# un intervalo con mas km/L que esto casi siempre es un odometro mal digitado
MAX_KM_PER_L = Decimal(60)
Q2 = Decimal("0.01")
Q3 = Decimal("0.001")


def d(v) -> Decimal:
    return Decimal(str(v)) if v is not None else Decimal(0)


def settings_of(tenant) -> dict:
    """Parametros del modulo dentro de tenant.settings["vehicles"] (no toca los ajustes generales)."""
    cfg = ((tenant.settings or {}).get("vehicles") or {}) if tenant else {}
    ref = cfg.get("monthly_reference")
    return {"monthly_reference": d(ref) if ref is not None else REFERENCE_DEFAULT}


# ---------- tipo de cambio: las cargas en dolares se analizan en colones ----------
class Fx:
    def __init__(self, db: Session):
        self.db = db
        self.cache: dict[tuple[str, date], Decimal | None] = {}

    def rate(self, currency: str, on: date) -> Decimal | None:
        if currency == "CRC":
            return Decimal(1)
        key = (currency, on)
        if key not in self.cache:
            fx = self.db.scalar(
                select(ExchangeRate).where(ExchangeRate.currency == currency, ExchangeRate.date <= on).order_by(ExchangeRate.date.desc()).limit(1)
            ) or self.db.scalar(select(ExchangeRate).where(ExchangeRate.currency == currency).order_by(ExchangeRate.date).limit(1))
            self.cache[key] = d(fx.sell) if fx else None
        return self.cache[key]

    def crc(self, amount, currency: str, on: date) -> Decimal | None:
        r = self.rate(currency or "CRC", on)
        return None if r is None else (d(amount) * r).quantize(Q2)


def exclusion(load: FuelLoad, vehicle: Vehicle | None) -> str | None:
    """Motivo por el que la carga no entra al analisis (None = entra)."""
    if vehicle is not None and vehicle.excluded_from_analysis:
        return "Vehículo excluido" + (f": {vehicle.exclusion_reason}" if vehicle.exclusion_reason else "")
    if load.excluded:
        return "Carga excluida" + (f": {load.exclusion_reason}" if load.exclusion_reason else "")
    return None


# ---------- km/L por intervalos entre tanques llenos ----------
@dataclass
class Interval:
    vehicle_id: int
    start: FuelLoad
    end: FuelLoad
    valid: bool
    reason: str | None = None
    km: int = 0
    liters: Decimal = Decimal(0)
    amount_crc: Decimal = Decimal(0)
    loads: list[int] = field(default_factory=list)

    @property
    def km_per_l(self) -> Decimal | None:
        return (Decimal(self.km) / self.liters).quantize(Q2) if self.valid and self.liters > 0 else None


def intervals(loads: list[FuelLoad], vehicle: Vehicle | None, fx: Fx) -> list[Interval]:
    """loads: TODAS las cargas del vehiculo (incluidas las excluidas) hasta el fin del periodo.

    Un intervalo va de un tanque lleno con odometro al siguiente. Es valido si el odometro sube, ninguna carga
    intermedia tiene un odometro fuera de orden, todas traen litros y ninguna esta excluida (si falta una carga
    los litros no cuadran). Los litros del intervalo son los cargados DESPUES del primer lleno hasta el segundo.
    """
    out: list[Interval] = []
    seq = sorted(loads, key=lambda x: (x.date, x.odometer if x.odometer is not None else -1, x.id))
    anchor_i = None
    doubtful = False  # el lleno anterior quedo con un odometro que retrocede: tampoco sirve de inicio
    for i, ld in enumerate(seq):
        # una carga excluida no sirve de ancla: su odometro puede ser de otro vehiculo
        if not (ld.full_tank and ld.odometer is not None and not ld.excluded):
            continue
        if anchor_i is None:
            anchor_i = i
            continue
        a = seq[anchor_i]
        between = seq[anchor_i + 1 : i + 1]
        it = Interval(vehicle_id=ld.vehicle_id, start=a, end=ld, valid=True, loads=[x.id for x in between])
        it.km = int(ld.odometer) - int(a.odometer)
        problems = []
        if doubtful:
            problems.append("el lleno inicial tiene un odómetro dudoso")
        doubtful = it.km <= 0
        if it.km <= 0:
            problems.append("el odómetro no sube entre los dos llenos")
        last = int(a.odometer)
        for x in between:
            if x.odometer is not None:
                if x.odometer < last or x.odometer > ld.odometer:
                    problems.append(f"odómetro fuera de orden el {x.date.isoformat()}")
                    break
                last = x.odometer
        if any(x.liters is None for x in between):
            problems.append("una carga intermedia no tiene litros")
        if any(x.excluded for x in between):
            problems.append("hay una carga excluida en medio")
        it.liters = sum((d(x.liters) for x in between), Decimal(0))
        amounts = [fx.crc(x.amount, x.currency, x.date) for x in between]
        if any(a_ is None for a_ in amounts):
            problems.append("falta tipo de cambio")
        it.amount_crc = sum((a_ for a_ in amounts if a_ is not None), Decimal(0))
        if not problems and it.liters > 0 and Decimal(it.km) / it.liters > MAX_KM_PER_L:
            problems.append("km/L imposible: revise el odómetro")
        if problems:
            it.valid, it.reason = False, "; ".join(problems)
        out.append(it)
        anchor_i = i
    return out


# ---------- meses (compras) ----------
def _months(a: date, b: date) -> list[tuple[int, int]]:
    out, y, m = [], a.year, a.month
    while (y, m) <= (b.year, b.month):
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def _partial_reason(y: int, m: int, a: date, b: date, first: date | None, last: date | None, today: date) -> str | None:
    m0, m1 = date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    if m0 <= today <= m1:
        return "mes en curso"
    if a > m0:
        return f"el periodo empieza el {a.strftime('%d/%m')}"
    if b < m1:
        return f"el periodo termina el {b.strftime('%d/%m')}"
    if first and m0 < first <= m1:
        return f"los datos empiezan el {first.strftime('%d/%m/%Y')}"
    if last and m0 <= last < m1:
        return f"los datos terminan el {last.strftime('%d/%m/%Y')}"
    return None


def report(db: Session, tenant, a: date, b: date, vehicle_id: int | None = None) -> dict:
    tid = tenant.id
    fx = Fx(db)
    vq = select(Vehicle).where(Vehicle.tenant_id == tid)
    if vehicle_id:
        vq = vq.where(Vehicle.id == vehicle_id)
    vehicles = {v.id: v for v in db.scalars(vq.order_by(Vehicle.name)).all()}
    if not vehicles:
        loads_all: list[FuelLoad] = []
    else:
        loads_all = list(
            db.scalars(select(FuelLoad).where(FuelLoad.tenant_id == tid, FuelLoad.vehicle_id.in_(vehicles), FuelLoad.date <= b).order_by(FuelLoad.date, FuelLoad.id)).all()
        )
    in_range = [x for x in loads_all if x.date >= a]
    included, excluded = [], []
    warnings: list[str] = []
    no_fx = 0
    for x in in_range:
        why = exclusion(x, vehicles.get(x.vehicle_id))
        crc = fx.crc(x.amount, x.currency, x.date)
        if crc is None:
            no_fx += 1
        (excluded if why else included).append((x, crc, why))
    if no_fx:
        warnings.append(f"{no_fx} carga(s) en otra moneda sin tipo de cambio registrado: no suman en colones.")

    # cobertura real de los datos (todas las cargas del tenant, no solo las del filtro)
    span = db.execute(select(func.min(FuelLoad.date), func.max(FuelLoad.date)).where(FuelLoad.tenant_id == tid)).one()
    first, last = span[0], span[1]
    today = date.today()

    # --- compras por mes ---
    buckets: dict[tuple[int, int], dict] = {}
    for y, m in _months(a, b):
        buckets[(y, m)] = {"month": f"{y:04d}-{m:02d}", "loads": 0, "liters": Decimal(0), "amount": Decimal(0), "liters_missing": 0}
    for x, crc, _ in included:
        bk = buckets[(x.date.year, x.date.month)]
        bk["loads"] += 1
        bk["amount"] += crc or 0
        if x.liters is None:
            bk["liters_missing"] += 1
        else:
            bk["liters"] += d(x.liters)
    months = []
    for (y, m), bk in buckets.items():
        reason = _partial_reason(y, m, a, b, first, last, today)
        months.append(
            {
                **bk,
                "liters": bk["liters"].quantize(Q3),
                "cost_per_liter": _cpl(included, y, m),
                "partial": reason is not None,
                "partial_reason": reason,
            }
        )

    # --- por vehiculo ---
    by_vehicle = []
    all_intervals: list[Interval] = []
    total_amount = sum((crc or 0 for _, crc, _ in included), Decimal(0))
    costs = db.execute(
        select(VehicleCost.vehicle_id, VehicleCost.kind, VehicleCost.amount, VehicleCost.currency, VehicleCost.date).where(
            VehicleCost.tenant_id == tid, VehicleCost.date >= a, VehicleCost.date <= b
        )
    ).all()
    other_costs: dict[int, dict[str, Decimal]] = defaultdict(lambda: defaultdict(Decimal))
    for vid, kind, amt, cur, on in costs:
        other_costs[vid][kind] += fx.crc(amt, cur, on) or 0
    for vid, v in vehicles.items():
        mine = [(x, crc) for x, crc, _ in included if x.vehicle_id == vid]
        ivs = [it for it in intervals([x for x in loads_all if x.vehicle_id == vid], v, fx) if a <= it.end.date <= b]
        all_intervals.extend(ivs)
        valid = [it for it in ivs if it.valid]
        km = sum(it.km for it in valid)
        lit = sum((it.liters for it in valid), Decimal(0))
        amt_iv = sum((it.amount_crc for it in valid), Decimal(0))
        amount = sum((crc or 0 for _, crc in mine), Decimal(0))
        liters = sum((d(x.liters) for x, _ in mine if x.liters is not None), Decimal(0))
        n_all = sum(1 for x in in_range if x.vehicle_id == vid)
        if not n_all and not v.active and not other_costs.get(vid):
            continue
        oc = other_costs.get(vid, {})
        by_vehicle.append(
            {
                "vehicle_id": vid,
                "name": v.name,
                "plate": v.plate,
                "owner": v.owner,
                "excluded": v.excluded_from_analysis,
                "exclusion_reason": v.exclusion_reason,
                "loads": len(mine),
                "loads_excluded": n_all - len(mine),
                "amount_excluded": sum((crc or 0 for x, crc, _ in excluded if x.vehicle_id == vid), Decimal(0)).quantize(Q2),
                "liters": liters.quantize(Q3),
                "amount": amount.quantize(Q2),
                "share_pct": float((amount / total_amount * 100).quantize(Q2)) if total_amount and not v.excluded_from_analysis else None,
                "cost_per_liter": (amount_l / liters).quantize(Q2) if (amount_l := _with_liters(mine)) and liters else None,
                "intervals": len(ivs),
                "intervals_valid": len(valid),
                "km_valid": km,
                "km_per_l": (Decimal(km) / lit).quantize(Q2) if km and lit else None,
                "cost_per_km": (amt_iv / km).quantize(Q2) if km else None,
                "other_costs": {k: oc[k].quantize(Q2) for k in COST_KINDS if oc.get(k)},
                "other_costs_total": sum(oc.values(), Decimal(0)).quantize(Q2),
            }
        )

    inc_loads = [x for x, _, _ in included]
    lit_tot = sum((d(x.liters) for x in inc_loads if x.liters is not None), Decimal(0))
    amt_with_l = _with_liters([(x, crc) for x, crc, _ in included])
    valid_all = [it for it in all_intervals if it.valid and not (vehicles.get(it.vehicle_id) and vehicles[it.vehicle_id].excluded_from_analysis)]
    km_tot = sum(it.km for it in valid_all)
    lit_iv = sum((it.liters for it in valid_all), Decimal(0))
    amt_iv = sum((it.amount_crc for it in valid_all), Decimal(0))
    missing_l = sum(1 for x in inc_loads if x.liters is None)
    if missing_l:
        warnings.append(f"{missing_l} carga(s) sin litros: suman en compras pero no en costo por litro ni km/L.")
    invalid = [it for it in all_intervals if not it.valid]
    if invalid:
        warnings.append(f"{len(invalid)} intervalo(s) entre llenos descartados para km/L (ver detalle).")
    if not valid_all:
        warnings.append("Sin intervalos válidos entre tanques llenos: km/L y costo por km no se pueden calcular todavía.")

    cfg = settings_of(tenant)
    n_months = len([m for m in months if not m["partial"]])
    return {
        "from": a.isoformat(),
        "to": b.isoformat(),
        "note": NOTE,
        "reference": {"amount": cfg["monthly_reference"], "label": REFERENCE_LABEL, "validated": False},
        "coverage": {"first": first.isoformat() if first else None, "last": last.isoformat() if last else None},
        "totals": {
            "loads": len(included),
            "liters": lit_tot.quantize(Q3),
            "amount": total_amount.quantize(Q2),
            "cost_per_liter": (amt_with_l / lit_tot).quantize(Q2) if lit_tot else None,
            "km_valid": km_tot,
            "km_per_l": (Decimal(km_tot) / lit_iv).quantize(Q2) if km_tot and lit_iv else None,
            "cost_per_km": (amt_iv / km_tot).quantize(Q2) if km_tot else None,
            "full_months": n_months,
            "excluded_loads": len(excluded),
            "excluded_amount": sum((crc or 0 for _, crc, _ in excluded), Decimal(0)).quantize(Q2),
        },
        "months": months,
        "by_vehicle": by_vehicle,
        "intervals": [
            {
                "vehicle_id": it.vehicle_id,
                "vehicle": vehicles[it.vehicle_id].name if it.vehicle_id in vehicles else "—",
                "from": it.start.date.isoformat(),
                "to": it.end.date.isoformat(),
                "odo_from": it.start.odometer,
                "odo_to": it.end.odometer,
                "km": it.km,
                "liters": it.liters.quantize(Q3),
                "amount": it.amount_crc.quantize(Q2),
                "km_per_l": it.km_per_l,
                "valid": it.valid,
                "reason": it.reason,
            }
            for it in sorted(all_intervals, key=lambda i: (i.end.date, i.vehicle_id))
        ],
        "excluded": [
            {"id": x.id, "date": x.date.isoformat(), "vehicle": vehicles[x.vehicle_id].name, "amount": d(x.amount), "currency": x.currency, "reason": why}
            for x, _, why in excluded
        ],
        "warnings": warnings,
    }


def _with_liters(pairs) -> Decimal:
    """Monto en colones solo de las cargas con litros (para que el costo por litro no se infle)."""
    return sum((crc or 0 for x, crc in pairs if x.liters is not None), Decimal(0))


def _cpl(included, y: int, m: int) -> Decimal | None:
    pairs = [(x, crc) for x, crc, _ in included if x.date.year == y and x.date.month == m]
    lit = sum((d(x.liters) for x, _ in pairs if x.liters is not None), Decimal(0))
    return (_with_liters(pairs) / lit).quantize(Q2) if lit else None


# ---------- importacion CSV / XLSX ----------
COLS = {
    "date": ("fecha", "dia", "date", "fecha_carga", "fecha_de_carga"),
    "vehicle": ("vehiculo", "placa", "auto", "carro", "unidad", "vehicle", "alias"),
    "driver": ("conductor", "chofer", "responsable", "driver", "funcionario"),
    "liters": ("litros", "lts", "cantidad_litros", "liters", "volumen"),
    "amount": ("importe", "monto", "total", "monto_total", "amount", "costo", "valor"),
    "currency": ("moneda", "divisa", "currency"),
    "odometer": ("odometro", "kilometraje", "km", "odometer", "kilometros"),
    "full_tank": ("tanque_lleno", "lleno", "tanque", "full_tank", "tipo_carga"),
    "usage": ("uso", "motivo", "destino", "detalle", "descripcion"),
    "project": ("proyecto", "numero_proyecto", "project"),
    "notes": ("notas", "observaciones", "comentarios"),
}
TEMPLATE = ["Fecha", "Vehiculo", "Conductor", "Litros", "Importe", "Moneda", "Odometro", "Tanque lleno", "Uso", "Proyecto", "Notas"]


def _cr(v, dec: int = 0) -> str:
    """Numero al estilo de Costa Rica: espacio de miles y coma decimal (30 600,00)."""
    s = f"{Decimal(str(v)):,.{dec}f}"
    return s.replace(",", " ").replace(".", ",")


def _yes(v) -> bool | None:
    s = tb.norm(v)
    if not s:
        return None
    if s in ("si", "s", "x", "1", "true", "lleno", "full", "yes", "tanque_lleno"):
        return True
    if s in ("no", "n", "0", "false", "parcial", "partial"):
        return False
    return None


def find_vehicle(vehicles: list[Vehicle], raw) -> Vehicle | None:
    s = tb.norm(raw)
    if not s:
        return None
    for v in vehicles:
        if v.plate and tb.norm(v.plate) == s:
            return v
    for v in vehicles:
        if tb.norm(v.name) == s:
            return v
    return None


def import_rows(db: Session, tid: int, rows: list[dict], commit: bool, user_id: int | None, create_vehicles: bool = False) -> list[dict]:
    """Vista previa (commit=False) o aplicacion. Acciones: nuevo | omitir | error; con avisos aparte.

    - Duplicado = misma fecha + vehiculo + importe (contra la base y dentro del mismo archivo) -> omitir.
    - Odometro que retrocede respecto de la carga anterior del vehiculo -> se importa con aviso (y el analisis
      descarta ese intervalo para km/L).
    - Faltan fecha, vehiculo o importe -> error. Faltan litros u odometro -> aviso.
    """
    vehicles = list(db.scalars(select(Vehicle).where(Vehicle.tenant_id == tid)).all())
    existing = defaultdict(list)  # vehicle_id -> [(date, odo)]
    dup_keys = set()
    for x in db.scalars(select(FuelLoad).where(FuelLoad.tenant_id == tid)).all():
        dup_keys.add((x.vehicle_id, x.date, d(x.amount).quantize(Q2)))
        if x.odometer is not None:
            existing[x.vehicle_id].append((x.date, x.odometer))
    projects = {tb.norm(p.number): p for p in db.scalars(select(Project).where(Project.tenant_id == tid)).all()}
    new_vehicles: dict[str, Vehicle] = {}
    parsed = []
    for i, r in enumerate(rows, start=2):
        v = {k: tb.pick(r, *al) for k, al in COLS.items()}
        errs, warns = [], []
        on = tb.to_date(v["date"])
        if not on:
            errs.append("fecha faltante o ilegible")
        elif on > date.today() + timedelta(days=1):
            errs.append("fecha en el futuro")
        veh = find_vehicle(vehicles, v["vehicle"])
        vname = tb.text(v["vehicle"], 120)
        if not vname:
            errs.append("vehículo faltante")
        elif not veh:
            if create_vehicles:
                key = tb.norm(vname)
                if key not in new_vehicles:
                    new_vehicles[key] = Vehicle(tenant_id=tid, name=vname, owner="empresa", active=True)
                veh = new_vehicles[key]
                warns.append(f"vehículo nuevo: {vname}")
            else:
                errs.append(f"vehículo no registrado: {vname}")
        amount = tb.num(v["amount"])
        if amount is None:
            errs.append("importe faltante")
        elif amount <= 0:
            errs.append("importe debe ser mayor que cero")
        liters = tb.num(v["liters"])
        if liters is None:
            warns.append("sin litros")
        elif liters <= 0:
            errs.append("litros deben ser mayores que cero")
        odo = tb.num(v["odometer"])
        if odo is None:
            warns.append("sin odómetro")
        full = _yes(v["full_tank"])
        cur = "USD" if tb.norm(v["currency"]) in ("usd", "dolares", "dolar", "us") else "CRC"
        proj = None
        if v["project"] not in (None, ""):
            proj = projects.get(tb.norm(v["project"]))
            if not proj:
                warns.append(f"proyecto {v['project']} no existe: se deja sin proyecto")
        parsed.append(
            {
                "row": i,
                "errs": errs,
                "warns": warns,
                "veh": veh,
                "vname": veh.name if veh else vname,
                "date": on,
                "amount": amount.quantize(Q2) if amount is not None else None,
                "liters": liters,
                "odo": int(odo) if odo is not None else None,
                "full": bool(full),
                "currency": cur,
                "driver": tb.text(v["driver"], 120) or (veh.driver if veh else None),
                "usage": tb.text(v["usage"], 160),
                "project": proj,
                "notes": tb.text(v["notes"], 2000),
            }
        )

    # odometros: se comparan contra la carga anterior del mismo vehiculo (base + archivo), por fecha
    timeline = defaultdict(list)
    for p in parsed:
        if not p["errs"] and p["odo"] is not None:
            timeline[id(p["veh"])].append(p)
    for _key, items in timeline.items():
        veh = items[0]["veh"]
        hist = existing.get(veh.id, []) if veh.id else []
        for p in items:
            prev = [o for (dt, o) in hist if dt < p["date"]] + [q["odo"] for q in items if q is not p and (q["date"], q["row"]) < (p["date"], p["row"])]
            later = [o for (dt, o) in hist if dt > p["date"]]
            if prev and p["odo"] < max(prev):
                p["warns"].append(f"odómetro retrocede ({_cr(p['odo'])} < {_cr(max(prev))})")
            elif later and p["odo"] > min(later):
                p["warns"].append("odómetro mayor que el de una carga posterior")

    out = []
    seen = set()
    for p in parsed:
        if p["errs"]:
            out.append({"row": p["row"], "action": "error", "detail": "; ".join(p["errs"]), "warnings": p["warns"], "vehicle": p["vname"]})
            continue
        key = (p["veh"].id or id(p["veh"]), p["date"], p["amount"])
        if (p["veh"].id and key in dup_keys) or key in seen:
            out.append({"row": p["row"], "action": "omitir", "detail": "duplicado: misma fecha, vehículo e importe", "warnings": p["warns"], "vehicle": p["vname"]})
            continue
        seen.add(key)
        out.append(
            {
                "row": p["row"],
                "action": "nuevo",
                "detail": f"{p['date'].strftime('%d/%m/%Y')} · {p['currency']} {_cr(p['amount'], 2)}" + (f" · {_cr(p['liters'], 2)} L" if p["liters"] is not None else ""),
                "warnings": p["warns"],
                "vehicle": p["vname"],
            }
        )
        if commit:
            veh = p["veh"]
            if not veh.id:
                db.add(veh)
                db.flush()
            db.add(
                FuelLoad(
                    tenant_id=tid,
                    vehicle_id=veh.id,
                    date=p["date"],
                    driver=p["driver"],
                    liters=p["liters"],
                    amount=p["amount"],
                    currency=p["currency"],
                    odometer=p["odo"],
                    full_tank=p["full"],
                    usage=p["usage"],
                    project_id=p["project"].id if p["project"] else None,
                    source="importacion",
                    notes=p["notes"],
                    created_by=user_id,
                )
            )
    return out


# ---------- Excel ----------
def to_xlsx(rep: dict, loads: list[dict], tenant_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    def sheet(ws, title: str, columns: list[str], rows: list[list]):
        ws.title = title[:30]
        ws["A1"] = f"{tenant_name} · {title}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = f"Periodo {rep['from']} a {rep['to']} · {NOTE}"
        ws["A2"].font = Font(color="8A858F")
        ws.append([])
        ws.append(columns)
        for c in ws[4]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="15131A")
        for r in rows:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
        for i, col in enumerate(columns, 1):
            width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:200]])
            ws.column_dimensions[get_column_letter(i)].width = min(48, width + 3)
        ws.freeze_panes = "A5"

    t = rep["totals"]
    sheet(
        wb.active,
        "Compras por mes",
        ["Mes", "Cargas", "Litros", "Compras (CRC)", "Costo por litro", "Parcial", "Motivo"],
        [[m["month"], m["loads"], m["liters"], m["amount"], m["cost_per_liter"], "sí" if m["partial"] else "no", m["partial_reason"] or ""] for m in rep["months"]]
        + [["Total", t["loads"], t["liters"], t["amount"], t["cost_per_liter"], "", ""], [REFERENCE_LABEL, "", "", rep["reference"]["amount"], "", "", "parámetro, no es un promedio"]],
    )
    sheet(
        wb.create_sheet(),
        "Por vehiculo",
        ["Vehículo", "Placa", "Propietario", "Excluido", "Motivo", "Cargas", "Litros", "Compras (CRC)", "% del total", "Costo por litro", "Km válidos", "Km/L", "Costo por km", "Otros costos"],
        [
            [
                v["name"], v["plate"] or "", v["owner"], "sí" if v["excluded"] else "no", v["exclusion_reason"] or "", v["loads"], v["liters"], v["amount"],
                v["share_pct"], v["cost_per_liter"], v["km_valid"], v["km_per_l"], v["cost_per_km"], v["other_costs_total"],
            ]
            for v in rep["by_vehicle"]
        ],
    )
    sheet(
        wb.create_sheet(),
        "Intervalos km-L",
        ["Vehículo", "Desde", "Hasta", "Odómetro inicial", "Odómetro final", "Km", "Litros", "Monto (CRC)", "Km/L", "Válido", "Motivo"],
        [[i["vehicle"], i["from"], i["to"], i["odo_from"], i["odo_to"], i["km"], i["liters"], i["amount"], i["km_per_l"], "sí" if i["valid"] else "no", i["reason"] or ""] for i in rep["intervals"]],
    )
    sheet(
        wb.create_sheet(),
        "Cargas",
        ["Fecha", "Vehículo", "Conductor", "Litros", "Importe", "Moneda", "Odómetro", "Tanque", "Uso", "Proyecto", "En análisis", "Motivo exclusión", "Origen"],
        [
            [
                x["date"], x["vehicle"], x["driver"] or "", x["liters"], x["amount"], x["currency"], x["odometer"], "lleno" if x["full_tank"] else "parcial",
                x["usage"] or "", x["project"] or "", "no" if x["excluded_effective"] else "sí", x["exclusion_effective"] or "", x["source"],
            ]
            for x in loads
        ],
    )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
