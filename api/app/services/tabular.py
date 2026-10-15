"""Lectura tolerante de CSV / Excel exportados por bancos, Fygaro u hojas propias.

- Encabezados normalizados (sin tildes, minusculas, espacios -> _) y resueltos por sinonimos.
- Numeros en formato CR/US ("1.234,56", "1,234.56", "(1 000)", "₡ 5 000").
- Fechas dd/mm/aaaa, aaaa-mm-dd, dd-mm-aa y fechas nativas de Excel.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation


class TableTooLarge(ValueError):
    """Hoja con mas filas, columnas o XML descomprimido de lo que se procesa (el mensaje se le muestra al usuario)."""


class TooManyRows(TableTooLarge):
    pass


# Topes de una hoja de calculo. Un .xlsx de 1,5 KB puede declarar <dimension ref="A1:XFD1048576"> y openpyxl (solo
# lectura) rellena filas y columnas hasta ahi: cientos de GB de None y el contenedor muere por falta de memoria. Por
# eso se revisa el zip antes de abrirlo, se ignora la dimension declarada y se recorre con topes explicitos.
MAX_GRID_ROWS = 50_000  # filas con datos (cada importador aplica ademas su propio tope, p. ej. 5000)
MAX_FILE_ROWS = 200_000  # numero de fila del archivo hasta donde se mira (filas vacias incluidas)
MAX_COLS = 128
XLSX_MAX_PART = 40 * 1024 * 1024  # cada parte del zip (hoja, textos compartidos) descomprimida
XLSX_MAX_TOTAL = 120 * 1024 * 1024
XLSX_MAX_RATIO = 100  # relacion de compresion de una parte grande (> 1 MB)
XLSX_MAX_ENTRIES = 500


def norm(s: object) -> str:
    t = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def _delimiter(text: str) -> str:
    """El separador que aparece la misma cantidad de veces en mas lineas (los bancos ponen titulos sin separador)."""
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    best, score = ",", -1
    for cand in (";", ",", "\t", "|"):
        counts = [ln.count(cand) for ln in lines if ln.count(cand) > 0]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        sc = counts.count(mode) * 100 + mode
        if sc > score:
            best, score = cand, sc
    return best


def read_table(data: bytes, filename: str, max_rows: int = MAX_GRID_ROWS) -> list[dict]:
    """Devuelve filas como dict {encabezado_normalizado: valor}. Salta filas vacias y titulos antes del encabezado."""
    grid = [r for _, r in read_grid(data, filename, max_rows)]
    if not grid:
        return []
    # el encabezado es la primera fila con al menos 2 celdas de texto (los bancos ponen titulos arriba)
    hi = next((i for i, r in enumerate(grid[:15]) if sum(1 for c in r if isinstance(c, str) and c.strip()) >= 2), 0)
    headers = [norm(h) or f"col_{i}" for i, h in enumerate(grid[hi])]
    rows = []
    for r in grid[hi + 1 :]:
        rows.append({headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))})
    return rows


def check_xlsx_zip(data: bytes) -> None:
    """Rechaza (TableTooLarge) un .xlsx cuyo zip declara partes enormes o con compresion sospechosa. zipfile corta la
    lectura en el tamano declarado, asi que estos topes acotan lo que openpyxl descomprime."""
    import zipfile

    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ValueError("el archivo de Excel está dañado") from e
    infos = zf.infolist()
    if len(infos) > XLSX_MAX_ENTRIES:
        raise TableTooLarge("el archivo de Excel tiene demasiadas partes")
    total = 0
    for i in infos:
        total += i.file_size
        big = i.file_size > 1024 * 1024
        if i.file_size > XLSX_MAX_PART or total > XLSX_MAX_TOTAL or (big and i.file_size / max(i.compress_size, 1) > XLSX_MAX_RATIO):
            raise TableTooLarge("la hoja de Excel es demasiado grande o sospechosa: divídala en partes o guárdela como CSV")


def _filled(r) -> bool:
    return any(c not in (None, "") and str(c).strip() for c in r)


def _xlsx_grid(data: bytes, max_rows: int) -> list[tuple[int, list]]:
    from openpyxl import load_workbook

    check_xlsx_zip(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.active
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()  # la dimension declarada no manda: un A1:XFD1048576 falso no rellena nada
        out: list[tuple[int, list]] = []
        for i, row in enumerate(ws.iter_rows(max_row=MAX_FILE_ROWS, max_col=MAX_COLS, values_only=True), start=1):
            if row.count(None) == len(row):
                continue  # fila vacia (o relleno hasta max_col): sin recorrerla en Python
            r = list(row)
            while r and r[-1] is None:
                r.pop()  # max_col rellena hasta MAX_COLS: se recorta lo vacio a la derecha
            if not _filled(r):
                continue
            out.append((i, r))
            if len(out) > max_rows:
                raise TooManyRows(f"máximo {max_rows} filas por archivo; divídalo en partes")
        return out
    finally:
        wb.close()


def read_grid(data: bytes, filename: str, max_rows: int = MAX_GRID_ROWS) -> list[tuple[int, list]]:
    """Celdas crudas sin filas vacias, con el numero de fila del archivo (1 = primera). Para formatos con
    encabezado mas abajo o filas de categoria (listas de proveedor), donde hay que mirar la grilla entera.
    Mas de `max_rows` filas con datos -> TableTooLarge."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:2] == b"PK":
        return _xlsx_grid(data, max_rows)
    else:
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        grid = [r for r in csv.reader(io.StringIO(text), delimiter=_delimiter(text))]
    out = [(i, r) for i, r in enumerate(grid, start=1) if _filled(r)]
    if len(out) > max_rows:
        raise TooManyRows(f"máximo {max_rows} filas por archivo; divídalo en partes")
    return out


def pick(row: dict, *aliases: str):
    """Primer valor no vacio entre los alias (ya normalizados)."""
    for a in aliases:
        v = row.get(a)
        if v not in (None, "") and str(v).strip() != "":
            return v
    return None


def num(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip()
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-") or s.endswith("-")
    s = re.sub(r"[^0-9,.\-]", "", s).strip("-")
    if not s:
        return None
    if "," in s and "." in s:
        # el ultimo separador es el decimal
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        parts = s.split(",")
        s = s.replace(",", ".") if len(parts) == 2 and len(parts[1]) <= 2 else s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    elif s.count(".") == 1 and len(s.split(".")[1]) == 3 and s.split(".")[0] not in ("", "0"):
        s = s.replace(".", "")  # "₡1.500" en montos: punto de miles
    try:
        n = Decimal(s)
    except InvalidOperation:
        return None
    return -n if neg else n


def to_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d-%m-%y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def text(v, limit: int = 300) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if isinstance(v, float) and v.is_integer():
        s = str(int(v))
    return s[:limit] or None
