"""Reportes (plan 3.6): cada uno devuelve columnas + filas; el router los sirve como JSON o Excel."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import Customer, Expense, Invoice, InvoiceLine, Payment, Product, StockMovement, Warehouse
from .inventory import stock_levels


@dataclass
class Report:
    key: str
    title: str
    columns: list[str]
    rows: list[list]
    totals: dict | None = None


def _cust(db: Session, cid: int | None) -> str:
    if not cid:
        return "—"
    c = db.get(Customer, cid)
    return c.name if c else "—"


def facturacion(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(
        select(Invoice)
        .where(Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada")
        .order_by(Invoice.issue_date, Invoice.id)
    ).all()
    return Report(
        "facturacion",
        "Facturación",
        ["Fecha", "Número", "Cliente", "Divisa", "Subtotal", "Descuento", "Impuesto", "Total", "Saldo", "Estado"],
        [
            [i.issue_date.isoformat(), i.number, _cust(db, i.customer_id), i.currency, i.subtotal, i.discount_total, i.tax_total, i.total, i.balance, i.status]
            for i in rows
        ],
        {"total": sum((Decimal(str(i.total)) for i in rows), Decimal(0)), "saldo": sum((Decimal(str(i.balance)) for i in rows), Decimal(0)), "n": len(rows)},
    )


def pendientes(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(
        select(Invoice)
        .where(Invoice.tenant_id == tid, Invoice.status.in_(("creado", "parcial", "vencida", "enviada")), Invoice.balance > 0)
        .order_by(Invoice.due_date)
    ).all()
    today = date.today()
    return Report(
        "pendientes",
        "Facturas pendientes",
        ["Número", "Cliente", "Emisión", "Vence", "Días vencida", "Divisa", "Total", "Saldo", "Estado"],
        [
            [
                i.number,
                _cust(db, i.customer_id),
                i.issue_date.isoformat(),
                i.due_date.isoformat() if i.due_date else "",
                max(0, (today - i.due_date).days) if i.due_date else 0,
                i.currency,
                i.total,
                i.balance,
                i.status,
            ]
            for i in rows
        ],
        {"saldo": sum((Decimal(str(i.balance)) for i in rows), Decimal(0)), "n": len(rows)},
    )


def transacciones(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(select(Payment).where(Payment.tenant_id == tid, Payment.paid_at >= a, Payment.paid_at <= b).order_by(Payment.paid_at, Payment.id)).all()
    inv = {i.id: i for i in db.scalars(select(Invoice).where(Invoice.id.in_([p.invoice_id for p in rows] or [0])))}
    return Report(
        "transacciones",
        "Transacciones",
        ["Fecha", "Factura", "Cliente", "Método", "Tipo", "Referencia", "Divisa", "Monto", "Propina", "Proveedor", "Estado"],
        [
            [
                p.paid_at.isoformat(),
                inv[p.invoice_id].number if p.invoice_id in inv else "",
                _cust(db, inv[p.invoice_id].customer_id) if p.invoice_id in inv else "",
                p.method,
                p.kind,
                p.external_ref or "",
                p.currency,
                p.amount,
                p.tip,
                p.provider,
                p.status,
            ]
            for p in rows
        ],
        {"capturado": sum((Decimal(str(p.amount)) for p in rows if p.kind == "captura" and p.status == "confirmado"), Decimal(0)), "n": len(rows)},
    )


def cierre_diario(db: Session, tid: int, a: date, b: date) -> Report:
    q = (
        select(Payment.paid_at, Payment.method, Payment.currency, func.sum(Payment.amount), func.count())
        .where(Payment.tenant_id == tid, Payment.paid_at >= a, Payment.paid_at <= b, Payment.kind == "captura", Payment.status == "confirmado")
        .group_by(Payment.paid_at, Payment.method, Payment.currency)
        .order_by(Payment.paid_at, Payment.method)
    )
    rows = [[r[0].isoformat(), r[1], r[2], r[3], r[4]] for r in db.execute(q)]
    return Report(
        "cierre", "Cierre diario", ["Fecha", "Método", "Divisa", "Monto", "Pagos"], rows, {"total": sum((Decimal(str(r[3])) for r in rows), Decimal(0))}
    )


def impuesto_facturado(db: Session, tid: int, a: date, b: date) -> Report:
    q = (
        select(InvoiceLine.tax_rate, func.sum(InvoiceLine.subtotal), func.sum(InvoiceLine.tax_amount), func.count())
        .join(Invoice, Invoice.id == InvoiceLine.invoice_id)
        .where(Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada")
        .group_by(InvoiceLine.tax_rate)
        .order_by(InvoiceLine.tax_rate.desc())
    )
    rows = [[f"{Decimal(str(r[0])).normalize()}%", r[1], r[2], r[3]] for r in db.execute(q)]
    return Report(
        "impuesto",
        "Impuesto facturado (por tarifa)",
        ["Tarifa", "Base imponible", "IVA", "Líneas"],
        rows,
        {"iva": sum((Decimal(str(r[2])) for r in rows), Decimal(0))},
    )


def iva(db: Session, tid: int, a: date, b: date) -> Report:
    deb = db.scalar(
        select(func.coalesce(func.sum(Invoice.tax_total), 0)).where(
            Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada"
        )
    )
    cred = db.scalar(
        select(func.coalesce(func.sum(Expense.tax_amount), 0)).where(
            Expense.tenant_id == tid, Expense.date >= a, Expense.date <= b, Expense.status != "anulado", Expense.iva_credit == "credito"
        )
    )
    prop = db.scalar(
        select(func.coalesce(func.sum(Expense.tax_amount), 0)).where(
            Expense.tenant_id == tid, Expense.date >= a, Expense.date <= b, Expense.status != "anulado", Expense.iva_credit == "proporcional"
        )
    )
    d, c, p = Decimal(str(deb)), Decimal(str(cred)), Decimal(str(prop))
    rows = [
        ["IVA débito (ventas)", d],
        ["IVA crédito (compras, genera crédito)", c],
        ["IVA compras proporcional (prorrata)", p],
        ["IVA a pagar (débito − crédito)", d - c],
    ]
    return Report("iva", "IVA · borrador D-104", ["Concepto", "Monto"], rows, {"a_pagar": d - c})


def gastos(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(
        select(Expense).where(Expense.tenant_id == tid, Expense.date >= a, Expense.date <= b, Expense.status != "anulado").order_by(Expense.date)
    ).all()
    return Report(
        "gastos",
        "Gastos",
        ["Fecha", "Descripción", "Divisa", "Subtotal", "IVA", "Total", "Crédito IVA", "Referencia", "Estado"],
        [[e.date.isoformat(), e.description, e.currency, e.subtotal, e.tax_amount, e.total, e.iva_credit, e.reference or "", e.status] for e in rows],
        {"total": sum((Decimal(str(e.total)) for e in rows), Decimal(0)), "n": len(rows)},
    )


def estado_resultados(db: Session, tid: int, a: date, b: date) -> Report:
    ventas = Decimal(
        str(
            db.scalar(
                select(func.coalesce(func.sum(Invoice.subtotal - Invoice.discount_total), 0)).where(
                    Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada"
                )
            )
        )
    )
    gastos_ = Decimal(
        str(
            db.scalar(
                select(func.coalesce(func.sum(Expense.subtotal), 0)).where(
                    Expense.tenant_id == tid, Expense.date >= a, Expense.date <= b, Expense.status != "anulado"
                )
            )
        )
    )
    return Report(
        "resultados",
        "Estado de resultados (sin IVA)",
        ["Concepto", "Monto"],
        [["Ventas netas", ventas], ["Gastos", gastos_], ["Resultado", ventas - gastos_]],
        {"resultado": ventas - gastos_},
    )


def inventario(db: Session, tid: int, a: date, b: date) -> Report:
    prods = {p.id: p for p in db.scalars(select(Product).where(Product.tenant_id == tid, Product.item_type == "producto"))}
    whs = {w.id: w for w in db.scalars(select(Warehouse).where(Warehouse.tenant_id == tid))}
    rows = []
    for lv in stock_levels(db, tid):
        p, w = prods.get(lv["product_id"]), whs.get(lv["warehouse_id"])
        if p:
            rows.append(
                [p.code, p.name, w.name if w else "", lv["quantity"], p.min_stock, p.price, p.currency, "BAJO" if lv["quantity"] <= p.min_stock else "OK"]
            )
    return Report("inventario", "Inventario", ["Código", "Producto", "Inventario", "Cantidad", "Mínimo", "Precio", "Divisa", "Alerta"], rows, {"n": len(rows)})


def venta_productos(db: Session, tid: int, a: date, b: date) -> Report:
    q = (
        select(InvoiceLine.code, InvoiceLine.name, func.sum(InvoiceLine.quantity), func.sum(InvoiceLine.subtotal), func.count(func.distinct(Invoice.id)))
        .join(Invoice, Invoice.id == InvoiceLine.invoice_id)
        .where(Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada")
        .group_by(InvoiceLine.code, InvoiceLine.name)
        .order_by(func.sum(InvoiceLine.subtotal).desc())
    )
    rows = [[r[0] or "", r[1], r[2], r[3], r[4]] for r in db.execute(q)]
    return Report("productos", "Venta de productos", ["Código", "Producto", "Cantidad", "Venta neta", "Facturas"], rows, {"n": len(rows)})


def movimientos(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(select(StockMovement).where(StockMovement.tenant_id == tid).order_by(StockMovement.at.desc()).limit(500)).all()
    prods = {p.id: p for p in db.scalars(select(Product).where(Product.tenant_id == tid))}
    whs = {w.id: w for w in db.scalars(select(Warehouse).where(Warehouse.tenant_id == tid))}
    return Report(
        "movimientos",
        "Movimientos de inventario",
        ["Fecha", "Producto", "Inventario", "Tipo", "Cantidad", "Referencia"],
        [
            [
                m.at.strftime("%Y-%m-%d %H:%M"),
                prods[m.product_id].name if m.product_id in prods else "",
                whs[m.warehouse_id].name if m.warehouse_id in whs else "",
                m.kind,
                m.quantity,
                m.reference or "",
            ]
            for m in rows
        ],
    )


# ---------- Sesion 4: ordenes, recepciones, propinas, D-151, planilla, conciliacion ----------
def _crc(amount, currency: str, fx) -> Decimal:
    v = Decimal(str(amount or 0))
    return v * Decimal(str(fx or 1)) if currency != "CRC" else v


def ordenes(db: Session, tid: int, a: date, b: date) -> Report:
    from datetime import datetime as dt

    from ..models import Order

    lo, hi = dt.combine(a, dt.min.time()), dt.combine(b, dt.max.time())
    rows = db.scalars(select(Order).where(Order.tenant_id == tid, Order.created_at >= lo, Order.created_at <= hi).order_by(Order.id)).all()
    return Report(
        "ordenes",
        "Órdenes",
        ["Fecha", "Orden", "Canal", "Cliente", "Envío", "Pago", "Estado", "Total", "Comprobante"],
        [
            [
                o.created_at.date().isoformat(),
                o.number,
                o.channel,
                (o.contact or {}).get("name", ""),
                o.shipping_method or "",
                o.payment_method or "",
                o.status,
                o.total,
                o.invoice_id or "",
            ]
            for o in rows
        ],
        {"Total": sum((Decimal(str(o.total)) for o in rows if o.status != "cancelado"), Decimal(0)), "Órdenes": len(rows)},
    )


def recepciones(db: Session, tid: int, a: date, b: date) -> Report:
    from ..models import ReceivedDocument

    rows = db.scalars(
        select(ReceivedDocument)
        .where(ReceivedDocument.tenant_id == tid, ReceivedDocument.issue_date >= a, ReceivedDocument.issue_date <= b)
        .order_by(ReceivedDocument.issue_date)
    ).all()
    return Report(
        "recepciones",
        "Recepciones",
        ["Fecha", "Emisor", "Cédula", "Consecutivo", "Moneda", "Subtotal", "IVA", "Total", "Condición IVA", "Respuesta", "Hacienda"],
        [
            [
                r.issue_date.isoformat() if r.issue_date else "",
                r.issuer_name or "",
                r.issuer_id or "",
                r.consecutive or "",
                r.currency,
                r.subtotal,
                r.tax_total,
                r.total,
                r.iva_condition,
                r.action or "pendiente",
                r.hacienda_status,
            ]
            for r in rows
        ],
        {
            "IVA": sum((Decimal(str(r.tax_total)) for r in rows if r.action != "rechazada"), Decimal(0)),
            "Total": sum((Decimal(str(r.total)) for r in rows if r.action != "rechazada"), Decimal(0)),
        },
    )


def propinas(db: Session, tid: int, a: date, b: date) -> Report:
    rows = db.scalars(
        select(Payment).where(Payment.tenant_id == tid, Payment.paid_at >= a, Payment.paid_at <= b, Payment.tip > 0).order_by(Payment.paid_at)
    ).all()
    invs = {i.id: i for i in db.scalars(select(Invoice).where(Invoice.id.in_({r.invoice_id for r in rows})))} if rows else {}
    return Report(
        "propinas",
        "Propinas",
        ["Fecha", "Factura", "Método", "Monto cobrado", "Propina"],
        [[r.paid_at.isoformat(), invs[r.invoice_id].number if r.invoice_id in invs else "", r.method, r.amount, r.tip] for r in rows],
        {"Propinas": sum((Decimal(str(r.tip)) for r in rows), Decimal(0))},
    )


D151_THRESHOLD = Decimal(2_500_000)
D151_SPECIFIC = Decimal(50_000)
SPECIFIC_WORDS = ("alquiler", "honorario", "comision", "servicios profesionales", "interes")


def d151(db: Session, tid: int, a: date, b: date) -> Report:
    """Borrador D-151 (resumen de clientes, proveedores y gastos especificos): montos anuales en colones por
    contraparte que superan el umbral. Revisar con el contador antes de presentar en TRIBU-CR."""
    from ..models import ExpenseCategory, ReceivedDocument, Supplier

    rows: list[list] = []
    sales: dict[tuple, Decimal] = {}
    for inv in db.scalars(select(Invoice).where(Invoice.tenant_id == tid, Invoice.issue_date >= a, Invoice.issue_date <= b, Invoice.status != "anulada")):
        if inv.doc_type not in ("FE", "FEE") or not inv.customer_id:
            continue
        c = db.get(Customer, inv.customer_id)
        if not c or not c.id_number:
            continue
        k = (c.id_number, c.name)
        sales[k] = sales.get(k, Decimal(0)) + _crc(inv.subtotal, inv.currency, inv.fx_sell)
    for (idn, name), amt in sorted(sales.items(), key=lambda x: -x[1]):
        if amt > D151_THRESHOLD:
            rows.append(["Ventas (clientes)", idn, name, amt.quantize(Decimal("0.01"))])

    buys: dict[tuple, Decimal] = {}
    specific: dict[tuple, Decimal] = {}
    cats = {c.id: c.name.lower() for c in db.scalars(select(ExpenseCategory).where(ExpenseCategory.tenant_id == tid))}
    for e in db.scalars(
        select(Expense).where(Expense.tenant_id == tid, Expense.date >= a, Expense.date <= b, Expense.status != "anulado", Expense.supplier_id.is_not(None))
    ):
        s = db.get(Supplier, e.supplier_id)
        if not s or not s.tax_id:
            continue
        k = (s.tax_id, s.name)
        amt = _crc(e.subtotal, e.currency, 1)
        cname = cats.get(e.category_id, "") + " " + e.description.lower()
        if any(w in cname for w in SPECIFIC_WORDS):
            specific[k] = specific.get(k, Decimal(0)) + amt
        else:
            buys[k] = buys.get(k, Decimal(0)) + amt
    for r in db.scalars(
        select(ReceivedDocument).where(
            ReceivedDocument.tenant_id == tid,
            ReceivedDocument.issue_date >= a,
            ReceivedDocument.issue_date <= b,
            ReceivedDocument.action.in_(("aceptada", "parcial")),
        )
    ):
        if r.expense_id or not r.issuer_id:  # si ya genero gasto con proveedor, se conto arriba
            continue
        k = (r.issuer_id, r.issuer_name or "")
        buys[k] = buys.get(k, Decimal(0)) + _crc(r.subtotal, r.currency, 1)
    for (idn, name), amt in sorted(buys.items(), key=lambda x: -x[1]):
        if amt > D151_THRESHOLD:
            rows.append(["Compras (proveedores)", idn, name, amt.quantize(Decimal("0.01"))])
    for (idn, name), amt in sorted(specific.items(), key=lambda x: -x[1]):
        if amt > D151_SPECIFIC:
            rows.append(["Gastos específicos", idn, name, amt.quantize(Decimal("0.01"))])
    return Report("d151", "D-151 (borrador)", ["Sección", "Identificación", "Nombre", "Monto anual ₡"], rows, {"Registros": len(rows)})


def planilla(db: Session, tid: int, a: date, b: date) -> Report:
    from ..models import PayrollRun

    runs = db.scalars(
        select(PayrollRun).where(PayrollRun.tenant_id == tid, PayrollRun.period_end >= a, PayrollRun.period_end <= b).order_by(PayrollRun.period_end)
    ).all()
    rows = []
    for r in runs:
        for ln in sorted(r.lines, key=lambda x: x.employee_name):
            rows.append(
                [
                    f"{r.period_start.isoformat()} a {r.period_end.isoformat()}",
                    r.status,
                    ln.employee_name,
                    ln.gross,
                    ln.ccss_worker,
                    ln.income_tax,
                    ln.other_deductions,
                    ln.net,
                    ln.ccss_employer,
                    ln.provisions,
                ]
            )
    live = [r for r in runs if r.status != "borrador"]
    return Report(
        "planilla",
        "Planilla",
        ["Periodo", "Estado", "Colaborador", "Bruto", "CCSS obrero", "Impuesto renta", "Otras ded.", "Neto", "CCSS patronal", "Provisiones"],
        rows,
        {
            "Bruto": sum((Decimal(str(r.gross)) for r in live), Decimal(0)),
            "Neto": sum((Decimal(str(r.net)) for r in live), Decimal(0)),
            "CCSS patronal": sum((Decimal(str(r.ccss_employer)) for r in live), Decimal(0)),
        },
    )


def conciliacion(db: Session, tid: int, a: date, b: date) -> Report:
    from ..routers.banking import summary

    rows = summary(db, tid, a, b)
    return Report(
        "conciliacion",
        "Conciliación bancaria",
        ["Fecha", "Cuenta", "Descripción", "Referencia", "Monto", "Estado", "Casado con", "Cómo"],
        rows,
        {"Pendientes": sum(1 for r in rows if r[5] == "pendiente"), "Conciliados": sum(1 for r in rows if r[5] == "conciliado")},
    )


def rentabilidad(db: Session, tid: int, a: date, b: date) -> Report:
    """En qué tipo de trabajo gana Crimson: venta, costo real y margen por proyecto cerrado."""
    from ..models import Project
    from ..routers.projects import economics
    from .documents import local_date, today_fx

    fx = today_fx(db, "USD")[0]
    rows, price_total, cost_total = [], Decimal(0), Decimal(0)
    cerrados = db.scalars(
        select(Project).where(Project.tenant_id == tid, Project.status.in_(("terminado", "entregado", "facturado"))).order_by(Project.id)
    ).all()
    for pr in cerrados:
        cierre = pr.end_date or local_date(pr.delivered_at) or local_date(pr.created_at)
        if not (a <= cierre <= b):
            continue
        e = economics(db, pr, fx)
        rows.append(
            [
                pr.number,
                pr.name,
                pr.status,
                cierre.isoformat(),
                e["price"],
                e["cost_materials"],
                e["cost_labor"],
                e["cost_travel"],
                e["cost_extra"],
                e["cost_real"],
                e["profit"],
                e["margin_real"],
                e["margin_planned"],
                e["hours"],
            ]
        )
        price_total += Decimal(str(e["price"] or 0))
        cost_total += Decimal(str(e["cost_real"]))
    rows.sort(key=lambda r: r[11])  # primero los que menos dejaron: ahi esta lo que hay que corregir
    return Report(
        "rentabilidad",
        "Rentabilidad por proyecto",
        [
            "Proyecto",
            "Nombre",
            "Estado",
            "Cierre",
            "Venta",
            "Equipos",
            "Mano de obra",
            "Viáticos",
            "Otros",
            "Costo real",
            "Utilidad",
            "Margen real %",
            "Margen cotizado %",
            "Horas",
        ],
        rows,
        {"Venta": price_total, "Costo real": cost_total, "Utilidad": price_total - cost_total},
    )


def rentabilidad_tipo(db: Session, tid: int, a: date, b: date) -> Report:
    """La pregunta de Andres: en que tipo de trabajo gana Crimson. CCTV puede dejar 24 % y redes 31 %,
    y hasta hoy eso solo se sabia por intuicion."""
    from ..models import Project
    from ..routers.projects import economics, project_solution
    from .documents import local_date, today_fx

    fx = today_fx(db, "USD")[0]
    por_tipo: dict[str, dict] = {}
    cerrados = db.scalars(select(Project).where(Project.tenant_id == tid, Project.status.in_(("terminado", "entregado", "facturado")))).all()
    for pr in cerrados:
        cierre = pr.end_date or local_date(pr.delivered_at) or local_date(pr.created_at)
        if not (a <= cierre <= b):
            continue
        e = economics(db, pr, fx)
        fila = por_tipo.setdefault(project_solution(db, pr), {"n": 0, "venta": Decimal(0), "costo": Decimal(0), "horas": 0.0})
        fila["n"] += 1
        fila["venta"] += Decimal(str(e["price"] or 0))
        fila["costo"] += Decimal(str(e["cost_real"]))
        fila["horas"] += e["hours"]
    rows = []
    for tipo, f in por_tipo.items():
        utilidad = f["venta"] - f["costo"]
        rows.append([tipo, f["n"], f["venta"], f["costo"], utilidad, _margen(f["venta"], f["costo"]), round(f["horas"], 1)])
    rows.sort(key=lambda r: -r[5])  # primero donde mas se gana
    venta = sum((Decimal(str(r[2])) for r in rows), Decimal(0))
    costo = sum((Decimal(str(r[3])) for r in rows), Decimal(0))
    return Report(
        "rentabilidad_tipo",
        "Rentabilidad por tipo de solución",
        ["Solución", "Proyectos", "Venta", "Costo real", "Utilidad", "Margen %", "Horas"],
        rows,
        {"Venta": venta, "Costo real": costo, "Utilidad": venta - costo, "Margen %": _margen(venta, costo)},
    )


def _margen(venta: Decimal, costo: Decimal) -> float:
    from ..services import pricing

    return pricing.margin_of(venta, costo)


REPORTS = {
    "facturacion": facturacion,
    "pendientes": pendientes,
    "transacciones": transacciones,
    "cierre": cierre_diario,
    "impuesto": impuesto_facturado,
    "iva": iva,
    "gastos": gastos,
    "resultados": estado_resultados,
    "inventario": inventario,
    "productos": venta_productos,
    "movimientos": movimientos,
    "ordenes": ordenes,
    "recepciones": recepciones,
    "propinas": propinas,
    "d151": d151,
    "planilla": planilla,
    "conciliacion": conciliacion,
    "rentabilidad": rentabilidad,
    "rentabilidad_tipo": rentabilidad_tipo,
}
CATALOG = [
    ("facturacion", "Facturación", "Todas las facturas del periodo con totales y saldo"),
    ("pendientes", "Facturas pendientes", "Por cobrar, con días de atraso"),
    ("impuesto", "Impuesto facturado", "Base e IVA por tarifa"),
    ("resultados", "Estado de resultados", "Ventas netas menos gastos"),
    ("cierre", "Cierre diario", "Cobros por día y método"),
    ("transacciones", "Transacciones", "Cada pago con método, tipo y referencia"),
    ("gastos", "Gastos", "Egresos registrados en contabilidad"),
    ("iva", "IVA", "Débito, crédito y prorrata · borrador D-104"),
    ("inventario", "Inventario", "Existencias por ubicación y alertas"),
    ("productos", "Venta de productos", "Ranking de productos y servicios"),
    ("movimientos", "Movimientos", "Ledger de inventario"),
    ("ordenes", "Órdenes", "Pedidos de tienda, POS y eventos"),
    ("recepciones", "Recepciones", "Facturas de proveedores recibidas y su respuesta"),
    ("propinas", "Propinas", "Propinas cobradas por pago"),
    ("d151", "D-151", "Borrador anual: clientes, proveedores y gastos específicos sobre el umbral"),
    ("planilla", "Planilla", "Salarios, cargas sociales e impuesto por colaborador"),
    ("conciliacion", "Conciliación bancaria", "Movimientos del banco y con qué se casaron"),
    ("rentabilidad", "Rentabilidad por proyecto", "Cuánto dejó cada instalación: venta contra costo real"),
    ("rentabilidad_tipo", "Rentabilidad por tipo de solución", "En qué tipo de trabajo gana Crimson: CCTV, redes, acceso…"),
]


def to_xlsx(rep: Report, a: date, b: date, tenant_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rep.title[:30]
    ws["A1"] = f"{tenant_name} · {rep.title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Periodo {a.isoformat()} a {b.isoformat()}"
    ws["A2"].font = Font(color="8A858F")
    ws.append([])
    ws.append(rep.columns)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="15131A")
        c.alignment = Alignment(vertical="center")
    for r in rep.rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in r])
    for i, col in enumerate(rep.columns, 1):
        width = max(len(str(col)), *(len(str(r[i - 1])) for r in rep.rows[:200])) if rep.rows else len(col)
        ws.column_dimensions[get_column_letter(i)].width = min(48, width + 3)
    ws.freeze_panes = "A5"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
