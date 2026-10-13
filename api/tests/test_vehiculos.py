"""Vehiculos y combustible (seccion 12 / AC21). Datos inventados: ningun nombre real."""

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from tests.test_session4 import _user_with_role


def _setup(client):
    pick = client.post("/vehicles", json={"name": "Pickup de prueba", "plate": "TST-101", "kind": "pickup", "owner": "empresa", "driver": "Conductor A"})
    assert pick.status_code == 201, pick.text
    # vehiculo personal que cargaba con la tarjeta: fuera del analisis SIN borrar su historia (AC21)
    suv = client.post(
        "/vehicles",
        json={"name": "SUV personal de prueba", "plate": "TST-202", "owner": "personal", "excluded_from_analysis": True, "exclusion_reason": "vehiculo personal"},
    )
    assert suv.status_code == 201, suv.text
    p, s = pick.json()["id"], suv.json()["id"]
    loads = [
        (p, "2026-07-01", 40, 30000, 10000, True),
        (p, "2026-07-10", 20, 15000, 10300, False),
        (p, "2026-07-20", 20, 15000, 10600, True),  # intervalo 1: 600 km / 40 L = 15 km/L
        (p, "2026-08-05", 30, 22500, 10500, True),  # odometro retrocede: intervalo invalido
        (p, "2026-08-20", 40, 30000, 11000, True),  # arranca de un lleno dudoso: invalido
        (s, "2026-07-15", 35, 50000, 5000, True),
        (s, "2026-08-15", 30, 45000, 5400, True),
    ]
    ids = []
    for vid, day, lit, amt, odo, full in loads:
        r = client.post("/vehicles/fuel", json={"vehicle_id": vid, "date": day, "liters": lit, "amount": amt, "odometer": odo, "full_tank": full})
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    # carga excluida a mano (con motivo): sigue en la lista
    r = client.post(
        "/vehicles/fuel",
        json={"vehicle_id": p, "date": "2026-08-10", "liters": 10, "amount": 7500, "odometer": 10800, "excluded": True, "exclusion_reason": "se cargo a otro vehiculo"},
    )
    assert r.status_code == 201
    return p, s, ids


def test_reporte_excluye_sin_borrar_y_kml_solo_con_llenos_coherentes(client, auth):
    p, s, _ = _setup(client)
    rep = client.get("/vehicles/report", params={"from": "2026-07-01", "to": "2026-08-31"}).json()
    t = rep["totals"]
    assert Decimal(str(t["amount"])) == Decimal(112500)  # solo el pickup, sin la carga excluida
    assert t["loads"] == 5 and t["excluded_loads"] == 3 and Decimal(str(t["excluded_amount"])) == Decimal(102500)
    assert Decimal(str(t["liters"])) == Decimal(150) and Decimal(str(t["cost_per_liter"])) == Decimal(750)
    assert t["km_valid"] == 600 and Decimal(str(t["km_per_l"])) == Decimal(15) and Decimal(str(t["cost_per_km"])) == Decimal(50)
    # meses: compras, no consumo; agosto parcial porque los datos terminan el 20/08
    assert rep["note"] == "Comprar en un mes no equivale a consumir en ese mes."
    jul, ago = rep["months"]
    assert Decimal(str(jul["amount"])) == Decimal(60000) and not jul["partial"]
    assert Decimal(str(ago["amount"])) == Decimal(52500) and ago["partial"] and "terminan" in ago["partial_reason"]
    # referencia: parametro no validado, nunca un promedio
    assert rep["reference"]["validated"] is False and Decimal(str(rep["reference"]["amount"])) == Decimal(200000)
    assert "no validada" in rep["reference"]["label"]
    # intervalos: uno valido, dos descartados con motivo
    ivs = [i for i in rep["intervals"] if i["vehicle_id"] == p]
    assert [i["valid"] for i in ivs] == [True, False, False]
    assert "no sube" in ivs[1]["reason"] and "dudoso" in ivs[2]["reason"]
    # el vehiculo excluido aparece con su marca pero no suma
    suv = next(v for v in rep["by_vehicle"] if v["vehicle_id"] == s)
    assert suv["excluded"] and suv["loads"] == 0 and suv["loads_excluded"] == 2 and suv["share_pct"] is None
    pick = next(v for v in rep["by_vehicle"] if v["vehicle_id"] == p)
    assert pick["share_pct"] == 100.0 and pick["km_valid"] == 600

    # la lista muestra todo, con la marca
    todas = client.get("/vehicles/fuel").json()
    assert len(todas) == 8
    excl = client.get("/vehicles/fuel", params={"status": "excluidas"}).json()
    assert len(excl) == 3 and all(x["exclusion_effective"] for x in excl)
    assert any(x["exclusion_effective"].startswith("Vehículo excluido") for x in excl)

    # periodo que corta el mes: julio parcial
    rep2 = client.get("/vehicles/report", params={"from": "2026-07-15", "to": "2026-07-31"}).json()
    assert rep2["months"][0]["partial"] and "empieza" in rep2["months"][0]["partial_reason"]


def test_reincluir_vehiculo_y_motivo_obligatorio(client, auth):
    p, s, _ = _setup(client)
    r = client.put(f"/vehicles/{s}", json={"name": "SUV personal de prueba", "owner": "personal", "excluded_from_analysis": True})
    assert r.status_code == 422  # excluir exige motivo
    r = client.put(f"/vehicles/{s}", json={"name": "SUV personal de prueba", "owner": "personal", "excluded_from_analysis": False})
    assert r.status_code == 200 and r.json()["exclusion_reason"] is None
    rep = client.get("/vehicles/report", params={"from": "2026-07-01", "to": "2026-08-31"}).json()
    assert Decimal(str(rep["totals"]["amount"])) == Decimal(112500 + 95000)
    assert client.get("/vehicles").json()[0]["loads"] >= 2


def test_config_referencia(client, auth):
    assert client.get("/vehicles/config").json()["monthly_reference"] in (200000, "200000")
    r = client.put("/vehicles/config", json={"monthly_reference": 180000})
    assert r.status_code == 200 and Decimal(str(r.json()["monthly_reference"])) == Decimal(180000)
    # los ajustes generales no se tocan (fuel_cost_per_km sigue igual) y no pisan la referencia
    assert client.get("/settings").json()["fuel_cost_per_km"] == 0
    client.put("/settings", json={"quote_valid_days": 20})
    assert Decimal(str(client.get("/vehicles/config").json()["monthly_reference"])) == Decimal(180000)


def _csv(rows: list[str]) -> bytes:
    return ("Fecha,Vehiculo,Conductor,Litros,Importe,Moneda,Odometro,Tanque lleno,Uso\n" + "\n".join(rows)).encode()


def test_importacion_vista_previa_valida_y_no_guarda_sin_confirmar(client, auth):
    _setup(client)
    data = _csv(
        [
            "20/07/2026,TST-101,Conductor A,20,15000,CRC,10600,si,Ruta",  # duplicado de la base
            "25/08/2026,TST-101,Conductor A,35,26250,CRC,11400,si,Visita",  # nueva
            "26/08/2026,Pickup de prueba,,10,7000,CRC,11100,no,Ruta",  # odometro retrocede
            "26/08/2026,Pickup de prueba,,10,7000,CRC,11100,no,Ruta",  # duplicado dentro del archivo
            ",TST-101,,10,7000,CRC,,no,",  # sin fecha
            "27/08/2026,ZZZ-999,,10,7000,CRC,,no,",  # vehiculo no registrado
            "28/08/2026,TST-101,,,9000,CRC,,no,",  # sin litros: aviso
        ]
    )
    antes = len(client.get("/vehicles/fuel").json())
    r = client.post("/vehicles/fuel/import?commit=false", files={"file": ("cargas.csv", data, "text/csv")})
    assert r.status_code == 200, r.text
    j = r.json()
    acts = [x["action"] for x in j["result"]]
    assert acts == ["omitir", "nuevo", "nuevo", "omitir", "error", "error", "nuevo"]
    assert "duplicado" in j["result"][0]["detail"]
    assert any("retrocede" in w for w in j["result"][2]["warnings"])
    assert any("sin litros" in w for w in j["result"][6]["warnings"])
    assert "no registrado" in j["result"][5]["detail"]
    assert j["columns"]["odometer"] == "odometro"
    assert len(client.get("/vehicles/fuel").json()) == antes  # la vista previa no guarda nada

    r = client.post("/vehicles/fuel/import?commit=true", files={"file": ("cargas.csv", data, "text/csv")})
    assert r.json()["summary"]["nuevo"] == 3
    nuevas = [x for x in client.get("/vehicles/fuel").json() if x["source"] == "importacion"]
    assert len(nuevas) == 3
    # reimportar el mismo archivo no duplica
    r = client.post("/vehicles/fuel/import?commit=true", files={"file": ("cargas.csv", data, "text/csv")})
    assert r.json()["summary"]["nuevo"] == 0
    tpl = client.get("/vehicles/fuel/import/template")
    assert tpl.status_code == 200 and tpl.content[:2] == b"PK"


def test_importacion_puede_crear_vehiculos(client, auth):
    data = _csv(["01/09/2026,Camion nuevo,,50,40000,CRC,1000,si,"])
    r = client.post("/vehicles/fuel/import?commit=false&create_vehicles=true", files={"file": ("c.csv", data, "text/csv")}).json()
    assert r["result"][0]["action"] == "nuevo" and "vehículo nuevo" in r["result"][0]["warnings"][0]
    assert client.get("/vehicles").json() == []
    client.post("/vehicles/fuel/import?commit=true&create_vehicles=true", files={"file": ("c.csv", data, "text/csv")})
    assert [v["name"] for v in client.get("/vehicles").json()] == ["Camion nuevo"]


def test_excel_y_otros_costos(client, auth):
    p, _, _ = _setup(client)
    r = client.post("/vehicles/costs", json={"vehicle_id": p, "date": "2026-08-01", "kind": "llantas", "amount": 120000, "description": "4 llantas"})
    assert r.status_code == 201
    rep = client.get("/vehicles/report", params={"from": "2026-07-01", "to": "2026-08-31"}).json()
    pick = next(v for v in rep["by_vehicle"] if v["vehicle_id"] == p)
    assert Decimal(str(pick["other_costs"]["llantas"])) == Decimal(120000)
    x = client.get("/vehicles/report", params={"from": "2026-07-01", "to": "2026-08-31", "format": "xlsx"})
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Compras por mes", "Por vehiculo", "Intervalos km-L", "Cargas"]
    assert wb["Cargas"].max_row == 4 + 8


def test_permisos_por_rol(client, auth, db_session):
    tid = auth["tenant"]["id"]
    _setup(client)
    admin_h = dict(client.headers)

    def login(email, role):
        pw = _user_with_role(db_session, tid, email, role)
        client.headers.pop("Authorization", None)
        tok = client.post("/auth/login", json={"email": email, "password": pw}).json()["access_token"]
        client.headers.update(admin_h)
        return {"Authorization": f"Bearer {tok}"}

    sup = login("supervisor-veh@ejemplo.com", "supervisor")
    assert client.get("/vehicles", headers=sup).status_code == 200
    assert client.get("/vehicles/report", headers=sup).status_code == 200
    assert client.post("/vehicles", json={"name": "X"}, headers=sup).status_code == 403
    assert client.get("/vehicles/report", params={"format": "xlsx"}, headers=sup).status_code == 403
    conta = login("conta-veh@ejemplo.com", "contabilidad")
    assert client.post("/vehicles", json={"name": "Sedan de prueba"}, headers=conta).status_code == 201
    assert client.get("/vehicles/report", params={"format": "xlsx"}, headers=conta).status_code == 200
    ven = login("ventas-veh@ejemplo.com", "ventas")
    assert client.get("/vehicles", headers=ven).status_code == 403
