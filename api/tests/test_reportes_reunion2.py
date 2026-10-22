"""Reunion 2 (21/10), pedidos 1 a 4: estado de resultados completo, comision pagada -> gasto, tablero de reportes,
Excel de venta de productos y respuesta sugerida a los comprobantes recibidos. Repo publico: solo datos ficticios."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import Commission, ExchangeRate, Expense, ExpenseCategory, Product, Project
from tests.aliados_helpers import interno, uid
from tests.hacienda_utils import FakeHacienda, clave, fe_xml, modo_pruebas

D = Decimal
A, B = "2026-05-01", "2026-05-31"
DIA = date(2026, 5, 12)


def _cust(client):
    return client.get("/customers").json()["items"][0]


def _prod(db_session, tid, code, kind):
    p = Product(tenant_id=tid, code=code, name=f"Artículo {code}", item_type=kind, price=0, currency="CRC")
    db_session.add(p)
    db_session.commit()
    return p.id


def _inv(client, lines, **extra):
    body = {"customer_id": _cust(client)["id"], "issue_date": DIA.isoformat(), "lines": lines, **extra}
    r = client.post("/invoices", json=body)
    assert r.status_code == 201, r.text
    return r.json()


def _tc(db_session, on=DIA, sell="500"):
    db_session.add(ExchangeRate(date=on, currency="USD", sell=D(sell), buy=D(sell) - 10))
    db_session.commit()


def _pnl(client, **params):
    r = client.get("/reports/resultados/completo", params={"from": A, "to": B, **params})
    assert r.status_code == 200, r.text
    return r.json()


def _vendedora(client, db_session, auth, email="vendedora.r2@ejemplo.com"):
    h = interno(client, db_session, auth["tenant"]["id"], email, "ventas")
    return h, uid(client, h)


def _comision_pendiente(client, db_session, auth, vid, monto=1000000):
    """Regla del 5 %, factura de la vendedora cobrada: nace una comision pendiente."""
    assert client.post("/commissions/rules", json={"user_id": vid, "name": "Vendedora", "percent": 5}).status_code == 201
    inv = _inv(client, [{"name": "Instalación", "quantity": 1, "unit_price": monto, "tax_rate": 13}], seller_id=vid)
    r = client.post(f"/invoices/{inv['id']}/payments", json={"method": "transferencia", "amount": monto * 1.13, "kind": "captura"})
    assert r.status_code == 201, r.text
    c = db_session.scalar(select(Commission).where(Commission.invoice_id == inv["id"]))
    assert c is not None and c.status in ("pendiente", "aprobada")
    return c.id, inv


# ---------------------------------------------------------------- 1. estado de resultados
def test_estado_de_resultados_por_tipo_vendedor_categoria_e_iva(client, auth, db_session):
    tid = auth["tenant"]["id"]
    _tc(db_session)
    eq, sv = _prod(db_session, tid, "R2-EQ", "producto"), _prod(db_session, tid, "R2-SV", "servicio")
    _vh, vid = _vendedora(client, db_session, auth)
    # factura mixta de la vendedora: 100 000 de equipo + 50 000 de servicio + 10 000 sin catalogo
    _inv(
        client,
        [
            {"product_id": eq, "name": "Cámara", "quantity": 2, "unit_price": 50000, "tax_rate": 13},
            {"product_id": sv, "name": "Configuración", "quantity": 1, "unit_price": 50000, "tax_rate": 13},
            {"name": "Viático", "quantity": 1, "unit_price": 10000, "tax_rate": 13},
        ],
        seller_id=vid,
    )
    # factura de proyecto (el proyecto la apunta): todo a Proyectos aunque tenga equipos
    pinv = _inv(client, [{"product_id": eq, "name": "Kit", "quantity": 1, "unit_price": 300000, "tax_rate": 13}])
    db_session.add(Project(tenant_id=tid, number="PRY-R2", name="Proyecto de prueba", invoice_id=pinv["id"]))
    db_session.commit()
    # factura en dolares con su tipo de cambio guardado (520): 100 USD = 52 000
    _inv(client, [{"product_id": sv, "name": "Soporte", "quantity": 1, "unit_price": 100, "tax_rate": 13}], currency="USD", fx_sell=520)
    # gastos: compra de equipo (costo de ventas por el nombre) y alquiler (operacion)
    compra = client.post("/expense-categories", json={"name": "Compras de equipo"}).json()["id"]
    alq = client.post("/expense-categories", json={"name": "Alquiler local"}).json()["id"]
    for cat, monto in ((compra, 80000), (alq, 40000)):
        r = client.post("/expenses", json={"description": "Gasto de prueba", "date": DIA.isoformat(), "subtotal": monto, "tax_rate": 13, "category_id": cat})
        assert r.status_code in (200, 201), r.text

    x = _pnl(client)
    tipos = {t["key"]: D(str(t["amount"])) for t in x["ingresos"]["por_tipo"]}
    assert tipos == {"equipos": D("100000.00"), "servicios": D("102000.00"), "proyectos": D("300000.00"), "otros": D("10000.00")}
    assert D(str(x["ingresos"]["total"])) == D("512000.00")
    vend = {v["seller_id"]: D(str(v["amount"])) for v in x["ingresos"]["por_vendedor"]}
    assert vend[vid] == D("160000.00") and sum(vend.values()) == D("512000.00")
    assert [c["label"] for c in x["costo_ventas"]["por_categoria"]] == ["Compras de equipo"]
    assert D(str(x["costo_ventas"]["total"])) == D("80000.00")  # con credito: el IVA no es gasto
    assert D(str(x["utilidad_bruta"])) == D("432000.00")
    assert D(str(x["utilidad_neta"])) == D("392000.00")
    assert D(str(x["iva"]["cobrado"])) == D("66560.00") and D(str(x["iva"]["pagado"])) == D("15600.00")
    # el grupo se puede fijar a mano: alquiler pasa a costo de ventas
    cats = {c["id"]: c for c in client.get("/expense-categories").json()}
    assert cats[alq]["pnl_group_effective"] == "operativo" and cats[alq]["pnl_group_auto"] is True
    r = client.put(f"/expense-categories/{alq}", json={"name": "Alquiler local", "active": True, "pnl_group": "costo_ventas"})
    assert r.status_code == 200 and r.json()["pnl_group"] == "costo_ventas"
    assert D(str(_pnl(client)["costo_ventas"]["total"])) == D("120000.00")
    # un cliente viejo que no manda pnl_group no lo borra
    client.put(f"/expense-categories/{alq}", json={"name": "Alquiler local", "active": True})
    assert db_session.get(ExpenseCategory, alq).pnl_group == "costo_ventas"

    # en dolares: cada documento con su tasa (la factura en USD queda en 100; lo demas / 500 del dia)
    u = _pnl(client, currency="USD")
    assert u["currency"] == "USD"
    tipos_usd = {t["key"]: D(str(t["amount"])) for t in u["ingresos"]["por_tipo"]}
    assert tipos_usd["servicios"] == D("200.00")  # 50 000 / 500 + 100 USD

    # vista generica (compatibilidad) y Excel con formato
    rep = client.get("/reports/resultados", params={"from": A, "to": B}).json()
    assert any(r[0] == "Utilidad bruta" for r in rep["rows"]) and any(r[0] == "Gastos" for r in rep["rows"])
    for path in ("/reports/resultados/completo", "/reports/resultados"):
        xl = client.get(path, params={"from": A, "to": B, "format": "xlsx"})
        assert xl.status_code == 200, xl.text
        wb = load_workbook(BytesIO(xl.content))
        assert wb.sheetnames == ["Estado de resultados", "Ventas por vendedor"]
        ws = wb["Estado de resultados"]
        textos = [ws.cell(i, 1).value for i in range(1, ws.max_row + 1)]
        assert "Utilidad bruta" in textos and "Total de ingresos" in textos
        fila = textos.index("Total de ingresos") + 1
        assert ws.cell(fila, 2).value == 512000 and "₡" in ws.cell(fila, 2).number_format


def test_estado_de_resultados_permisos(client, auth, db_session):
    tid = auth["tenant"]["id"]
    conta = interno(client, db_session, tid, "conta.r2@ejemplo.com", "contabilidad")
    ventas = interno(client, db_session, tid, "ventas.r2@ejemplo.com", "ventas")
    assert client.get("/reports/resultados/completo", headers=conta).status_code == 200
    assert client.get("/reports/resultados/completo", headers=ventas).status_code == 403
    assert client.get("/reports/resultados/completo", params={"currency": "EUR"}).status_code == 422


# ---------------------------------------------------------------- 2. comision pagada -> gasto
def test_comision_pagada_registra_un_gasto_una_vez_y_deshacer_lo_anula(client, auth, db_session):
    tid = auth["tenant"]["id"]
    _vh, vid = _vendedora(client, db_session, auth)
    cid, _inv_ = _comision_pendiente(client, db_session, auth, vid)
    r = client.patch("/commissions", json={"ids": [cid], "status": "pagada", "paid_on": DIA.isoformat(), "paid_ref": "Planilla mayo"})
    assert r.status_code == 200, r.text
    db_session.expire_all()
    c = db_session.get(Commission, cid)
    e = db_session.get(Expense, c.expense_id)
    assert e is not None and e.reference == f"COMISION-{cid}" and e.status == "pagado" and D(str(e.total)) == D("50000")
    assert db_session.get(ExpenseCategory, e.category_id).name == "Comisiones" and e.iva_credit == "no_credito" and e.date == DIA
    # marcarla otra vez no duplica
    assert client.patch("/commissions", json={"ids": [cid], "status": "pagada"}).status_code == 409
    assert db_session.scalars(select(Expense).where(Expense.reference == f"COMISION-{cid}")).all() == [e]
    # sale en el estado de resultados como gasto de operacion "Comisiones"
    x = _pnl(client)
    com = next(g for g in x["gastos_operacion"]["por_categoria"] if g["label"] == "Comisiones")
    assert com["commissions"] and D(str(com["amount"])) == D("50000.00") and D(str(x["comisiones"])) == D("50000.00")
    fila = next(r for r in client.get("/commissions").json()["rows"] if r["id"] == cid)
    assert fila["expense_id"] == e.id and fila["can_undo_payment"] is True

    # deshacer el pago: solo el admin; el gasto queda anulado (no se borra)
    conta = interno(client, db_session, tid, "conta.com@ejemplo.com", "contabilidad")
    assert client.post("/commissions/undo-payment", json={"ids": [cid], "reason": "Pago no salió"}, headers=conta).status_code == 403
    r = client.post("/commissions/undo-payment", json={"ids": [cid], "reason": "Pago no salió del banco"})
    assert r.status_code == 200 and r.json()["actualizadas"] == 1, r.text
    db_session.expire_all()
    assert db_session.get(Commission, cid).status == "aprobada" and db_session.get(Expense, e.id).status == "anulado"
    assert not any(g["label"] == "Comisiones" for g in _pnl(client)["gastos_operacion"]["por_categoria"])
    # volver a pagarla reutiliza el mismo gasto
    assert client.patch("/commissions", json={"ids": [cid], "status": "pagada", "paid_on": DIA.isoformat()}).status_code == 200
    db_session.expire_all()
    assert db_session.get(Commission, cid).expense_id == e.id and db_session.get(Expense, e.id).status == "pagado"
    assert len(db_session.scalars(select(Expense).where(Expense.reference == f"COMISION-{cid}")).all()) == 1


def test_comision_pagada_antes_se_registra_desde_resultados(client, auth, db_session):
    _vh, vid = _vendedora(client, db_session, auth, "vendedora.vieja@ejemplo.com")
    cid, _ = _comision_pendiente(client, db_session, auth, vid)
    client.patch("/commissions", json={"ids": [cid], "status": "pagada", "paid_on": DIA.isoformat()})
    # simula una pagada antes del registro automatico
    c = db_session.get(Commission, cid)
    e = db_session.get(Expense, c.expense_id)
    c.expense_id = None
    db_session.delete(e)
    db_session.commit()
    x = _pnl(client)
    assert x["comisiones_sin_gasto"] == 1 and x["can_register_commissions"] and x["avisos"]
    r = client.post("/commissions/paid-expenses", json={"from": A, "to": B})
    assert r.status_code == 200 and r.json()["registrados"] == 1
    assert _pnl(client)["comisiones_sin_gasto"] == 0
    assert client.post("/commissions/paid-expenses", json={"from": A, "to": B}).json()["registrados"] == 0


# ---------------------------------------------------------------- 3. tablero y venta de productos
def test_tablero_indicadores_y_margen_solo_admin(client, auth, db_session):
    tid = auth["tenant"]["id"]
    eq = _prod(db_session, tid, "R2-TB", "producto")
    _inv(client, [{"product_id": eq, "name": "Grabador", "quantity": 3, "unit_price": 100000, "tax_rate": 13}])
    t = client.get("/reports/tablero/indicadores", params={"from": A, "to": B})
    assert t.status_code == 200, t.text
    t = t.json()
    assert D(str(t["ventas"]["total"])) == D("300000.00") and t["ventas"]["facturas"] == 1
    assert len(t["serie"]) == 6 and t["serie"][-1]["month"] == 5
    assert t["top_productos"][0]["name"] == "Grabador" and D(str(t["top_productos"][0]["quantity"])) == 3
    assert {x["key"] for x in t["cxc"]["antiguedad"]} == {"al_dia", "d30", "d60", "d90", "mas90"}
    assert t["margen"] is not None and "margen_bruto_pct" in t["margen"]
    conta = interno(client, db_session, tid, "conta.tb@ejemplo.com", "contabilidad")
    tc = client.get("/reports/tablero/indicadores", params={"from": A, "to": B}, headers=conta).json()
    assert tc["margen"] is None and D(str(tc["ventas"]["total"])) == D("300000.00")
    ventas = interno(client, db_session, tid, "ventas.tb@ejemplo.com", "ventas")
    assert client.get("/reports/tablero/indicadores", headers=ventas).status_code == 403


def test_excel_venta_de_productos_completo(client, auth, db_session):
    tid = auth["tenant"]["id"]
    _vh, vid = _vendedora(client, db_session, auth, "vendedora.xl@ejemplo.com")
    eq = _prod(db_session, tid, "R2-XL", "producto")
    _inv(
        client,
        [{"product_id": eq, "name": "Disco duro", "quantity": 2, "unit_price": 50000, "discount_type": "percent", "discount_value": 10, "tax_rate": 13}],
        seller_id=vid,
    )
    rep = client.get("/reports/productos", params={"from": A, "to": B}).json()
    cols = rep["columns"]
    for c in ("Cliente", "Vendedor", "Cantidad", "Precio unitario", "Descuento", "Impuesto", "Total", "Moneda"):
        assert c in cols
    fila = next(r for r in rep["rows"] if r[cols.index("Producto")] == "Disco duro")
    assert D(str(fila[cols.index("Descuento")])) == D("10000.00") and D(str(fila[cols.index("Subtotal")])) == D("90000")
    assert fila[cols.index("Vendedor")] == "Ventas Prueba" and D(str(fila[cols.index("Total")])) == D("101700")
    xl = client.get("/reports/productos", params={"from": A, "to": B, "format": "xlsx"})
    wb = load_workbook(BytesIO(xl.content))
    assert wb.sheetnames[1] == "Resumen por producto"
    ws = wb.worksheets[0]
    head = [c.value for c in ws[4]]
    assert head == cols
    assert ws.cell(5, cols.index("Total") + 1).number_format.startswith("#,##0.00")
    pie = [ws.cell(i, 1).value for i in range(5, ws.max_row + 1)]
    assert "Total ₡" in pie and "Líneas" in pie


# ---------------------------------------------------------------- 4. respuesta sugerida a Hacienda
HOY = date.today()
FECHA = f"{HOY.isoformat()}T10:00:00-06:00"


def _cl(n, emisor="003101000001"):
    return clave(fecha=HOY.strftime("%d%m%y"), emisor=emisor, consecutivo=f"0010000101{n:010d}")


def _up(client, n, **kw):
    kw.setdefault("fecha", FECHA)
    kw.setdefault("consecutivo", f"0010000101{n:010d}")
    r = client.post("/reception/upload", files={"file": ("f.xml", fe_xml(_cl(n, kw.pop("emisor", "003101000001")), **kw).encode(), "application/xml")})
    assert r.status_code == 201, r.text
    return r.json()


def test_respuesta_sugerida_se_aprueba_con_un_clic(client, auth, db_session):
    d1 = _up(client, 71)
    det = client.get(f"/reception/{d1['id']}").json()
    p = det["proposal"]
    assert p["ready"] and p["action"] == "aceptada" and p["iva_condition"] == "credito" and p["source"] == "predeterminado" and p["activity_code"]
    lista = client.get("/reception/documents", params={"status": "por_revisar"}).json()["items"]
    assert next(i for i in lista if i["id"] == d1["id"])["proposal"]["ready"]
    s = client.get("/reception/summary").json()
    assert s["listas_para_aprobar"] >= 1 and s["buzon"]["state"] == "apagado" and "sin revisar" not in s["buzon"]["label"]
    # contabilidad responde la primera a mano: gasto corriente (04) con su categoria
    cat = client.post("/expense-categories", json={"name": "Suministros"}).json()["id"]
    r = client.post(f"/reception/{d1['id']}/respond", json={"action": "aceptada", "iva_condition": "no_credito", "category_id": cat})
    assert r.status_code == 200 and r.json()["status"] == "respondido"
    # la siguiente del mismo proveedor sugiere lo mismo y se aprueba con un clic (por el flujo de siempre)
    d2 = _up(client, 72)
    p2 = client.get(f"/reception/{d2['id']}").json()["proposal"]
    assert p2["source"] == "proveedor" and p2["iva_condition"] == "no_credito" and p2["category_id"] == cat
    r = client.post(f"/reception/{d2['id']}/approve")
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["status"] == "respondido" and out["action"] == "aceptada" and out["iva_condition"] == "no_credito"
    assert out["responses"][0]["consecutive"][8:10] == "05" and out["proposal"] is None
    e = db_session.get(Expense, out["expense_id"])
    assert e.category_id == cat and e.received_document_id == d2["id"]
    assert client.post(f"/reception/{d2['id']}/approve").status_code == 409  # ya respondido: no se responde dos veces
    # en lote
    d3, d4 = _up(client, 73), _up(client, 74)
    r = client.post("/reception/approve-batch", json={"ids": [d3["id"], d4["id"], d2["id"]]}).json()
    assert {x["id"] for x in r["aprobadas"]} == {d3["id"], d4["id"]} and [o["id"] for o in r["omitidas"]] == [d2["id"]]
    # el panel tributario dice el estado del buzon en palabras
    ts = client.get("/accounting/tax-summary", params={"month": HOY.strftime("%Y-%m")}).json()
    assert ts["hacienda"]["buzon"]["state"] == "apagado"


def test_sugerencia_no_lista_si_hay_un_gasto_parecido_y_lectura_no_aprueba(client, auth, db_session):
    d = _up(client, 75, emisor="003101000002")
    # un gasto manual del mismo monto, sin proveedor: hay que decidir si vincularlo
    client.post("/expenses", json={"description": "Internet", "date": HOY.isoformat(), "subtotal": 100000, "tax_rate": 13})
    p = client.get(f"/reception/{d['id']}").json()["proposal"]
    if client.get(f"/reception/{d['id']}").json()["expense_candidates"]:
        assert not p["ready"] and p["blockers"]
        assert client.post(f"/reception/{d['id']}/approve").status_code == 409
    ventas = interno(client, db_session, auth["tenant"]["id"], "ventas.rc@ejemplo.com", "ventas")
    assert client.post(f"/reception/{d['id']}/approve", headers=ventas).status_code == 403


def test_aprobar_en_pruebas_envia_a_hacienda(client, auth, db_session, monkeypatch, p12_empresa):
    fake = FakeHacienda(sequence=["procesando", "aceptado"]).install(monkeypatch)
    modo_pruebas(client, *p12_empresa)
    d = _up(client, 76)
    r = client.post(f"/reception/{d['id']}/approve")
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["responses"][0]["mode"] == "pruebas" and out["hacienda_status"] in ("enviado", "procesando", "aceptado")
    assert len(fake.posts) == 1 and fake.posts[0]["clave"] == _cl(76)
